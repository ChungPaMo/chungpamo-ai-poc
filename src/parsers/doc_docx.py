"""DOCX 직접 파서 (python-docx) → 공통 중간 JSON 0.2.

- 본문 요소(w:p, w:tbl)를 문서 순서대로 읽는다. body_index는 본문 w:p·w:tbl 순번(0부터).
- 병합: w:gridSpan(가로), w:vMerge restart/continue(세로)를 XML에서 직접 읽어 앵커 셀 + span으로 표현.
- 머리말·꼬리말은 part=header/footer 블록. 필드(쪽 번호)는 텍스트가 없으므로 존재 여부만 format_specific에 기록.
- 표 캡션·각주: 스타일로 구분되지 않으면 일반 문단(paragraph)으로 둔다 (추측하지 않음).
"""

from __future__ import annotations

from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.table import _Cell
from docx.text.paragraph import Paragraph

from .doc_common import diag, finalize_order, new_ir

LOC_BASE = {"body_index": 0, "table_index": 0, "section_index": 0, "paragraph_index": 0, "row": 1, "col": 1}


def _block_type(style_name: str) -> tuple[str, int | None]:
    s = (style_name or "").lower()
    if s == "title":
        return "heading", 0
    if s.startswith("heading"):
        digits = "".join(ch for ch in s if ch.isdigit())
        return "heading", int(digits) if digits else None
    if s == "caption":
        return "caption", None
    if s.startswith("list"):
        return "list_item", None
    return "paragraph", None


def _tc_props(tc) -> tuple[int, str | None]:
    pr = tc.tcPr
    span, vmerge = 1, None
    if pr is not None:
        gs = pr.find(qn("w:gridSpan"))
        if gs is not None:
            span = int(gs.get(qn("w:val")))
        vm = pr.find(qn("w:vMerge"))
        if vm is not None:
            vmerge = vm.get(qn("w:val")) or "continue"
    return span, vmerge


def _table_cells(tbl, doc, body_index: int, table_index: int) -> tuple[list[dict], int, int, list[int]]:
    cells, anchors_by_col = [], {}
    header_rows, n_cols = [], 0
    trs = tbl.findall(qn("w:tr"))
    for r, tr in enumerate(trs, start=1):
        trpr = tr.find(qn("w:trPr"))
        if trpr is not None and trpr.find(qn("w:tblHeader")) is not None:
            header_rows.append(r)
        c = 1
        for tc in tr.findall(qn("w:tc")):
            span, vmerge = _tc_props(tc)
            if vmerge == "continue" and c in anchors_by_col:
                anchors_by_col[c]["rowspan"] += 1
            else:
                text = "\n".join(p.text for p in _Cell(tc, doc).paragraphs)
                cell = {"row": r, "col": c, "rowspan": 1, "colspan": span, "text": text, "raw": text,
                        "is_header": r in header_rows,
                        "location": {"format": "docx", "part": "body", "body_index": body_index, "table_index": table_index, "row": r, "col": c}}
                cells.append(cell)
                if vmerge == "restart":
                    anchors_by_col[c] = cell
                else:
                    anchors_by_col.pop(c, None)
            c += span
        n_cols = max(n_cols, c - 1)
    return cells, len(trs), n_cols, header_rows


def parse_docx(path: Path, rel: str) -> dict:
    ir = new_ir(path, rel, "docx", {"name": "python-docx", "version": _dist_version("python-docx"), "options": {}})
    d = docx.Document(str(path))
    seq: list[str] = []
    nb = nt = 0

    def add_block(text, btype, level, loc, style):
        nonlocal nb
        nb += 1
        bid = f"b{nb:04d}"
        ir["blocks"].append({"block_id": bid, "type": btype, "order": None, "text": text, "location": loc, "level": level, "style": {"name": style}})
        return bid

    # 머리말 (읽기 순서상 본문 앞)
    footers = []
    page_field = []
    for si, sec in enumerate(d.sections):
        for part, hf in (("header", sec.header), ("footer", sec.footer)):
            if hf.is_linked_to_previous and si > 0:
                continue
            if "PAGE" in hf._element.xml:
                page_field.append({"part": part, "section_index": si})
            for pi, p in enumerate(hf.paragraphs):
                if p.text.strip():
                    loc = {"format": "docx", "part": part, "section_index": si, "paragraph_index": pi}
                    if part == "header":
                        seq.append(add_block(p.text, "header", None, loc, p.style.name))
                    else:
                        footers.append((p.text, loc, p.style.name))
    body = [c for c in d.element.body.iterchildren() if c.tag in (qn("w:p"), qn("w:tbl"))]
    for bi, el in enumerate(body):
        if el.tag == qn("w:p"):
            p = Paragraph(el, d)
            if not p.text.strip():
                continue
            btype, level = _block_type(p.style.name if p.style is not None else "")
            seq.append(add_block(p.text, btype, level, {"format": "docx", "part": "body", "body_index": bi}, p.style.name if p.style is not None else None))
        else:
            if el.find(".//" + qn("w:tbl")) is not None:
                diag(ir, "EXT-UNSUPPORTED", "중첩 표는 펼치지 않음", location={"format": "docx", "part": "body", "body_index": bi})
            cells, n_rows, n_cols, header_rows = _table_cells(el, d, bi, nt)
            tid = f"t{nt + 1:03d}"
            ir["tables"].append({"table_id": tid, "location": {"format": "docx", "part": "body", "body_index": bi, "table_index": nt},
                                 "n_rows": n_rows, "n_cols": n_cols, "header_rows": len(header_rows) or None, "cells": cells,
                                 "continuation": None, "caption_block_id": None, "detection": "native"})
            seq.append(tid)
            nt += 1
    for text, loc, style in footers:
        seq.append(add_block(text, "footer", None, loc, style))
    finalize_order(ir, seq)
    has_footnotes = any("footnotes" in str(r.reltype) for r in d.part.rels.values())
    if has_footnotes:
        diag(ir, "EXT-UNSUPPORTED", "python-docx는 각주(footnotes part) 읽기 API가 없음 — 각주 텍스트 미추출")
    ir["document"]["units_of_division"] = {"sections": len(d.sections)}
    ir["format_specific"].update({
        "location_index_base": LOC_BASE,
        "page_number_fields": page_field,
        "footnotes_part_present": has_footnotes,
        "pages": "DOCX에는 쪽 나눔 결과가 저장되지 않음 (뷰어 렌더링 의존) — page 위치 미제공",
        "caption_link": "표 캡션은 스타일이 일반 문단이면 구분하지 않음 (caption_block_id = null)",
    })
    return ir


def _dist_version(name: str) -> str:
    from importlib.metadata import version
    return version(name)
