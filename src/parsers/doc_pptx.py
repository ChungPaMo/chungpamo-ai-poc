"""PPTX 직접 파서 (python-pptx) → 공통 중간 JSON 0.2.

- 슬라이드(1부터) → 도형 트리 순서(z-order)로 읽는다. 그룹 도형은 재귀하며 nested_path에 바깥→안 도형 이름을 남긴다.
- 표 도형: is_merge_origin·span_height·span_width로 앵커 셀 + span, is_spanned 셀은 건너뜀.
- 발표자 노트: part=notes, type=speaker_notes (해당 슬라이드 도형 뒤에 둔다).
- 블록 유형: 제목 자리표시자만 slide_title, 그 밖의 텍스트는 slide_body (도형 이름으로 의미를 추측하지 않음).
- 위치 bbox는 EMU [left, top, width, height] (그룹 안 도형은 그룹 좌표계 원값).
"""

from __future__ import annotations

from importlib.metadata import version
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from .doc_common import finalize_order, link_fragments, new_ir

TITLE_PH = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
LOC_BASE = {"slide_index": 1, "row": 1, "col": 1}


def parse_pptx(path: Path, rel: str) -> dict:
    ir = new_ir(path, rel, "pptx", {"name": "python-pptx", "version": version("python-pptx"), "options": {}})
    prs = Presentation(str(path))
    seq: list[str] = []
    groups = []
    counters = {"b": 0, "t": 0}

    def loc_of(shape, si, slide_id, nested):
        ph = shape.placeholder_format.type if shape.is_placeholder else None
        return {"format": "pptx", "part": "slide", "slide_index": si, "slide_id": slide_id, "shape_id": shape.shape_id,
                "shape_name": shape.name, "placeholder_type": str(ph) if ph is not None else None,
                "bbox": [shape.left, shape.top, shape.width, shape.height], "bbox_unit": "EMU", "nested_path": nested}

    def walk(shapes, si, slide_id, path):
        for sh in shapes:
            nested = path + [sh.name]
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                groups.append({"slide_index": si, "group_shape_id": sh.shape_id, "group_name": sh.name,
                               "children": [{"shape_id": c.shape_id, "name": c.name} for c in sh.shapes]})
                walk(sh.shapes, si, slide_id, nested)
                continue
            if getattr(sh, "has_table", False) and sh.has_table:
                counters["t"] += 1
                tid = f"t{counters['t']:03d}"
                tbl = sh.table
                cells = []
                for r, row in enumerate(tbl.rows, start=1):
                    for c, cell in enumerate(row.cells, start=1):
                        if cell.is_spanned:
                            continue
                        rs, cs = (cell.span_height, cell.span_width) if cell.is_merge_origin else (1, 1)
                        cells.append({"row": r, "col": c, "rowspan": rs, "colspan": cs, "text": cell.text, "raw": cell.text,
                                      "is_header": bool(tbl.first_row) and r == 1,
                                      "location": {"format": "pptx", "part": "slide", "slide_index": si, "shape_id": sh.shape_id, "row": r, "col": c}})
                ir["tables"].append({"table_id": tid, "location": loc_of(sh, si, slide_id, nested), "n_rows": len(tbl.rows), "n_cols": len(tbl.columns),
                                     "header_rows": None, "cells": cells, "continuation": None, "caption_block_id": None, "detection": "native",
                                     "first_row_flag": bool(tbl.first_row)})
                seq.append(tid)
                continue
            if sh.has_text_frame and sh.text_frame.text.strip():
                counters["b"] += 1
                bid = f"b{counters['b']:04d}"
                ph = sh.placeholder_format.type if sh.is_placeholder else None
                ir["blocks"].append({"block_id": bid, "type": "slide_title" if ph in TITLE_PH else "slide_body", "order": None,
                                     "text": sh.text_frame.text, "location": loc_of(sh, si, slide_id, nested), "level": None})
                seq.append(bid)

    for si, slide in enumerate(prs.slides, start=1):
        walk(slide.shapes, si, slide.slide_id, [])
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None and slide.notes_slide.notes_text_frame.text.strip():
            counters["b"] += 1
            bid = f"b{counters['b']:04d}"
            ir["blocks"].append({"block_id": bid, "type": "speaker_notes", "order": None, "text": slide.notes_slide.notes_text_frame.text,
                                 "location": {"format": "pptx", "part": "notes", "slide_index": si, "slide_id": slide.slide_id, "nested_path": ["notes"]},
                                 "level": None})
            seq.append(bid)
    finalize_order(ir, seq)
    links = link_fragments(ir, "slide_index")
    ir["document"]["units_of_division"] = {"slides": len(prs.slides)}
    ir["format_specific"].update({
        "location_index_base": LOC_BASE,
        "slide_size_emu": [prs.slide_width, prs.slide_height],
        "layouts": [s.slide_layout.name for s in prs.slides],
        "groups": groups,
        "reading_order_basis": "도형 트리 순서(z-order), 노트는 슬라이드 끝",
        "continuation_links": {"count": links, "method": "header_repeat_rule (python-pptx 자체 기능 아님)"},
    })
    return ir
