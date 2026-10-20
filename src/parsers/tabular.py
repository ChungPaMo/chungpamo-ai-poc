"""E02 표 형식 파일 읽기: 인코딩 판별, CSV 구분자·헤더, Excel 시트·병합 헤더.

표준 csv 모듈과 openpyxl만 사용한다. 원본 파일은 읽기만 한다.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

ENCODING_CANDIDATES = ("utf-8", "cp949")
DELIMITER_CANDIDATES = (",", ";", "\t", "|")


@dataclass
class Table:
    """읽은 표. rows는 물리 행 순서 그대로(헤더 포함), row 번호는 1부터."""
    path: str
    sheet: str
    rows: list[list]
    header_rows: int
    header: list[dict]                       # [{col_index(1부터), group, name}]
    meta: dict = field(default_factory=dict)

    def data_rows(self):
        for n, r in enumerate(self.rows[self.header_rows:], start=self.header_rows + 1):
            if any(v not in (None, "") for v in r):
                yield n, r


def detect_encoding(data: bytes) -> tuple[str, bool]:
    """(인코딩, BOM 여부). BOM이 있으면 utf-8-sig, 아니면 UTF-8 → CP949 순으로 엄격 해독."""
    if data.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig", True
    for enc in ENCODING_CANDIDATES:
        try:
            data.decode(enc)
            return enc, False
        except UnicodeDecodeError:
            continue
    raise UnicodeDecodeError("unknown", data[:1], 0, 1, "지원 인코딩(UTF-8, CP949)으로 해독 불가")


def detect_delimiter(text: str) -> str:
    """모든 행에서 같은 열 수(>1)로 나뉘는 후보. 여러 개면 열 수가 가장 많은 후보."""
    best = None
    for cand in DELIMITER_CANDIDATES:
        counts = {len(r) for r in csv.reader(io.StringIO(text, newline=""), delimiter=cand) if r}
        if len(counts) == 1:
            n = counts.pop()
            if n > 1 and (best is None or n > best[1]):
                best = (cand, n)
    if best is None:
        raise ValueError("행마다 열 수가 같은 구분자를 찾지 못함")
    return best[0]


def read_csv(path: Path, rel: str) -> Table:
    data = Path(path).read_bytes()
    enc, bom = detect_encoding(data)
    text = data.decode(enc)
    delim = detect_delimiter(text)
    rows = [r for r in csv.reader(io.StringIO(text, newline=""), delimiter=delim)]
    newline = "CRLF" if b"\r\n" in data else "LF"
    header = [{"col_index": i, "group": None, "name": h} for i, h in enumerate(rows[0], start=1)]
    return Table(rel, "", rows, 1, header, {"format": "csv", "encoding": enc, "bom": bom, "delimiter": delim, "newline": newline})


def read_xlsx(path: Path, rel: str, sheet: str) -> Table:
    """1행에 병합 범위가 있으면 2행 헤더(1행 = 그룹, 2행 = 열 이름)로 복원한다."""
    wb = openpyxl.load_workbook(path, read_only=False, data_only=True)
    ws = wb[sheet]
    rows = [[c.value for c in row] for row in ws.iter_rows()]
    merged_row1 = [m for m in ws.merged_cells.ranges if m.min_row == 1 and m.max_row == 1]
    if merged_row1:
        header_rows = 2
        group_at = {}
        for c, v in enumerate(rows[0], start=1):
            if v not in (None, ""):
                group_at[c] = str(v)
        for m in merged_row1:
            g = rows[0][m.min_col - 1]
            for c in range(m.min_col, m.max_col + 1):
                group_at[c] = str(g)
        names = rows[1]
        header = [{"col_index": c, "group": group_at.get(c), "name": None if v is None else str(v)}
                  for c, v in enumerate(names, start=1)]
    else:
        header_rows = 1
        header = [{"col_index": c, "group": None, "name": None if v is None else str(v)} for c, v in enumerate(rows[0], start=1)]
    wb.close()
    return Table(rel, sheet, rows, header_rows, header,
                 {"format": "xlsx", "merged_header_ranges": [str(m) for m in merged_row1], "header_rows": header_rows})
