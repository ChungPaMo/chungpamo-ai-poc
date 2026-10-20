"""PPTX 직접 파서 v2 (python-pptx) → 공통 중간 JSON 0.2.

v1 대비 추가 (외부 문서 Baseline에서 확인한 실패 유형):
1. 차트: python-pptx 차트 API로 제목·계열 이름·항목 이름을 블록(type=other, format_specific kind=chart)으로 출력
2. SmartArt: 그래픽 프레임의 다이어그램 데이터 부품(dgm:dataModel)에서 노드(node·asst) 텍스트를 읽어 블록으로 출력
   (python-pptx에 SmartArt API 없음 → 관계로 부품을 찾아 XML 직접 읽기)
그 밖(도형 순서·그룹·표·노트·표 조각 연결 규칙)은 v1과 같다.
"""

from __future__ import annotations

from importlib.metadata import version

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from ..doc_common import diag, finalize_order, link_fragments, new_ir
from . import IMPL_VERSION

TITLE_PH = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
LOC_BASE = {"slide_index": 1, "row": 1, "col": 1}
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
DGM = "{http://schemas.openxmlformats.org/drawingml/2006/diagram}"
R_DM = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}dm"


def _txbody_lines(tx) -> str:
    paras = []
    for p in tx.findall(A + "p"):
        paras.append("".join((e.findtext(A + "t") or "") if e.tag in (A + "r", A + "fld") else "\n" if e.tag == A + "br" else "" for e in p))
    return "\n".join(paras)


def chart_text(chart) -> list[str]:
    out = []
    try:
        if chart.has_title and chart.chart_title.has_text_frame:
            out.append(chart.chart_title.text_frame.text)
    except Exception:
        pass
    for plot in chart.plots:
        try:
            out += [str(c) for c in plot.categories]
        except Exception:
            pass
        for s in plot.series:
            if s.name:
                out.append(s.name)
    seen, uniq = set(), []
    for s in out:
        if s and s.strip() and s not in seen:
            seen.add(s)
            uniq.append(s)
    return uniq


def smartart_text(shape) -> list[str] | None:
    gd = shape._element.find(f".//{A}graphicData")
    rel = gd.find(f"{DGM}relIds") if gd is not None else None
    if rel is None or rel.get(R_DM) is None:
        return None
    try:
        part = shape.part.related_part(rel.get(R_DM))
    except KeyError:
        return None
    root = etree.fromstring(part.blob)
    out = []
    for pt in root.iter(DGM + "pt"):
        if pt.get("type", "node") in ("node", "asst"):
            t = pt.find(DGM + "t")
            if t is not None and _txbody_lines(t).strip():
                out.append(_txbody_lines(t))
    return out


def parse_pptx_v2(path, rel: str) -> dict:
    ir = new_ir(path, rel, "pptx", {"name": "python-pptx", "version": version("python-pptx"), "impl_version": IMPL_VERSION, "options": {}})
    prs = Presentation(str(path))
    seq, groups = [], []
    counters = {"b": 0, "t": 0}
    extras = {"charts": 0, "smartart": 0}

    def loc_of(shape, si, slide_id, nested):
        ph = shape.placeholder_format.type if shape.is_placeholder else None
        return {"format": "pptx", "part": "slide", "slide_index": si, "slide_id": slide_id, "shape_id": shape.shape_id,
                "shape_name": shape.name, "placeholder_type": str(ph) if ph is not None else None,
                "bbox": [shape.left, shape.top, shape.width, shape.height], "bbox_unit": "EMU", "nested_path": nested}

    def add_block(text, btype, loc, kind=None):
        counters["b"] += 1
        bid = f"b{counters['b']:04d}"
        b = {"block_id": bid, "type": btype, "order": None, "text": text, "location": loc, "level": None}
        if kind:
            b["format_specific"] = {"kind": kind}
        ir["blocks"].append(b)
        seq.append(bid)

    def walk(shapes, si, slide_id, path):
        for sh in shapes:
            nested = path + [sh.name]
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                groups.append({"slide_index": si, "group_shape_id": sh.shape_id, "group_name": sh.name,
                               "children": [{"shape_id": c.shape_id, "name": c.name} for c in sh.shapes]})
                walk(sh.shapes, si, slide_id, nested)
                continue
            if getattr(sh, "has_table", False) and sh.has_table:
                counters["t"] += 1
                tid = f"t{counters['t']:03d}"
                tbl, cells = sh.table, []
                for r, row in enumerate(tbl.rows, start=1):
                    for c, cell in enumerate(row.cells, start=1):
                        if cell.is_spanned:
                            continue
                        rs, cs = (cell.span_height, cell.span_width) if cell.is_merge_origin else (1, 1)
                        cells.append({"row": r, "col": c, "rowspan": rs, "colspan": cs, "text": cell.text, "raw": cell.text,
                                      "is_header": bool(tbl.first_row) and r == 1,
                                      "location": {"format": "pptx", "part": "slide", "slide_index": si, "shape_id": sh.shape_id, "row": r, "col": c}})
                ir["tables"].append({"table_id": tid, "location": loc_of(sh, si, slide_id, nested), "n_rows": len(tbl.rows), "n_cols": len(tbl.columns),
                                     "header_rows": None, "cells": cells, "continuation": None, "caption_block_id": None, "detection": "native",
                                     "first_row_flag": bool(tbl.first_row)})
                seq.append(tid)
                continue
            if getattr(sh, "has_chart", False) and sh.has_chart:
                texts = chart_text(sh.chart)
                if texts:
                    extras["charts"] += 1
                    add_block("\n".join(texts), "other", loc_of(sh, si, slide_id, nested), "chart")
                continue
            if sh._element.tag.endswith("}graphicFrame"):
                nodes = smartart_text(sh)
                if nodes:
                    extras["smartart"] += 1
                    add_block("\n".join(nodes), "other", loc_of(sh, si, slide_id, nested), "smartart")
                elif nodes is None:
                    diag(ir, "EXT-UNSUPPORTED", "표·차트·SmartArt가 아닌 그래픽 프레임 내용은 읽지 않음", location={"slide_index": si, "shape_id": sh.shape_id})
                continue
            if sh.has_text_frame and sh.text_frame.text.strip():
                ph = sh.placeholder_format.type if sh.is_placeholder else None
                add_block(sh.text_frame.text, "slide_title" if ph in TITLE_PH else "slide_body", loc_of(sh, si, slide_id, nested))

    for si, slide in enumerate(prs.slides, start=1):
        walk(slide.shapes, si, slide.slide_id, [])
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None and slide.notes_slide.notes_text_frame.text.strip():
            add_block(slide.notes_slide.notes_text_frame.text, "speaker_notes",
                      {"format": "pptx", "part": "notes", "slide_index": si, "slide_id": slide.slide_id, "nested_path": ["notes"]})
    finalize_order(ir, seq)
    links = link_fragments(ir, "slide_index")
    ir["document"]["units_of_division"] = {"slides": len(prs.slides)}
    ir["format_specific"].update({
        "location_index_base": LOC_BASE, "slide_size_emu": [prs.slide_width, prs.slide_height],
        "layouts": [s.slide_layout.name for s in prs.slides], "groups": groups,
        "reading_order_basis": "도형 트리 순서(z-order), 노트는 슬라이드 끝",
        "continuation_links": {"count": links, "method": "header_repeat_rule (python-pptx 자체 기능 아님)"},
        "graphic_frames_read": extras,
    })
    return ir
