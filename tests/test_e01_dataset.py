"""E01 데이터셋·정답셋 테스트.

- 손 계산 대조: 핵심 수치는 아래 주석의 산식을 사람이 직접 계산한 값으로 고정했다 (구현 A·B 어느 쪽 함수도 쓰지 않음).
- 음성 대조: 값을 일부러 바꾸면 독립 검증(구현 B)이 실패를 보고하는지 확인한다.
- UCI 원본(data/raw, Git 미추적)이 없으면 원본 의존 테스트는 사유를 남기고 건너뛴다.
"""

from __future__ import annotations

import ast
import csv
import hashlib
import io
import json
import shutil
import subprocess
import tempfile
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

import e01_verify
from e01_common import PROJECT_ROOT, Paths, read_json, sha256_file, write_json

ROOT = PROJECT_ROOT
P = Paths()
needs_uci = pytest.mark.skipif(not P.uci_csv.exists(), reason="UCI 원본 없음: scripts/e01_download_uci.py 실행 필요")


@pytest.fixture(scope="module")
def gt():
    return read_json(P.ground_truth)


@pytest.fixture(scope="module")
def X(gt):
    return gt["expected"]


def d(x) -> Decimal:
    return Decimal(str(x))


# ---------------------------------------------------------------------------
# 손 계산 대조
# ---------------------------------------------------------------------------
def test_production_requirements_hand(X):
    r = {p: d(v["net_requirement"]) for p, v in X["production_requirements"].items()}
    # 완제품: OPEN 주문 합 + 안전재고 − (보유 − 할당)
    assert r["FG001"] == (80 + 100 + 90 + 110 + 60 + 70) + 50 - (120 - 20)          # 460
    assert r["FG002"] == (40 + 60 + 50 + 70 + 30) + 30 - (60 - 0)                    # 220 (O018 취소 제외)
    assert r["FG003"] == (120 + 150 + 130 + 160 + 80 + 100 + 40) + 40 - (90 - 10)    # 740
    # 반제품: 부모 순소요를 모두 합한 뒤 재고를 한 번만 차감
    assert r["SF001"] == (460 * 1 + 220 * 2) + 20 - 150                               # 770
    assert r["SF002"] == (740 * 1) + 10 - 40                                          # 710


def test_semi_finished_netted_once(X):
    sf = X["production_requirements"]["SF001"]
    assert d(sf["dependent_demand"]) == 900 and d(sf["available"]) == 150
    # 부모별로 따로 차감했다면 (460-150)+(440-150)=600+20 이 되었을 것 — 그 값이 아님을 확인
    assert d(sf["net_requirement"]) != (460 - 150) + (440 - 150) + 20


def test_material_and_shortage_hand(X):
    g = {m: d(v) for m, v in X["material_gross_requirements"].items()}
    s = {m: d(v["shortage"]) for m, v in X["material_shortage"].items()}
    assert g["M001"] == Decimal("460") * Decimal("1.2") + Decimal("740") * Decimal("0.6")       # 996 kg (FG001 V1, FG003 V2)
    assert s["M001"] == Decimal("996") + 50 - (800 + 150)                                     # 96 (입고예정 11-20 포함)
    assert g["M003"] == Decimal("770") * Decimal("0.3")                                         # 300 g → 0.3 kg, 231 kg
    assert s["M003"] == Decimal("231") + 20 - Decimal("200000") * Decimal("0.001")              # 재고 g 기록 → 51 kg
    assert g["M002"] == Decimal("220") * Decimal("0.8") + Decimal("710") * Decimal("0.4") * Decimal("1.02")  # 465.68 (손실률 2%)
    assert s["M002"] == 0                                                                       # 465.68+30-500 < 0 (경계)
    assert g["M009"] == Decimal("740") * Decimal("0.06")                                        # V2 60 g → 44.4 kg
    assert s["M009"] == Decimal("4.4")
    assert s["M007"] == 710 + 20 - 400 and s["M006"] == 460 + 50 - 500 and s["M010"] == (460 + 220) + 50 - 700


def test_bom_selected_by_quantity(X):
    assert X["production_requirements"]["FG003"]["bom_id"] == "BOM-FG003-V2"   # 740 ≥ 500
    assert X["production_requirements"]["FG001"]["bom_id"] == "BOM-FG001-V1"   # 기준일 2026-11-16 ≤ 11-30


def test_bom_version_boundaries(X):
    got = {(c["product_id"], c["date"], c["qty"]): (c["status"], c["candidates"]) for c in X["cq"]["CQ-04"]}
    assert got[("FG001", "2026-11-30", "460")] == ("SELECTED", ["BOM-FG001-V1"])   # 종료일 포함
    assert got[("FG001", "2026-12-01", "460")] == ("SELECTED", ["BOM-FG001-V2"])
    assert got[("FG003", "2026-11-16", "499")] == ("SELECTED", ["BOM-FG003-V1"])   # 상한 미포함
    assert got[("FG003", "2026-11-16", "500")] == ("SELECTED", ["BOM-FG003-V2"])
    assert got[("FG002", "2025-12-31", "10")] == ("NONE", [])                      # 유효 시작 전 → 선택 불가


def test_order_o005_hand(X):
    r = next(o for o in X["order_feasibility"] if o["order_id"] == "O005")
    # FG002 가용 60: O002가 40 사용 → 20 남음 → O005 60 중 20 재고, 40 생산
    assert d(r["from_stock"]) == 20 and d(r["production_qty"]) == 40
    # SF001 가용 150: O013 40, O004 100 사용 → 10 남음 → 필요 80 중 70 생산 → M003(부족, LT 12일) 대기
    assert d(r["semi_finished_production"]["SF001"]) == 70
    assert r["material_ready_date"] == "2026-11-27"                 # 2026-11-15 + 12일
    assert r["production_days"] == 3 + 2                            # FG002 3일 + SF001 2일 (근무일)
    assert r["expected_completion_date"] == "2026-12-04"           # 11-27(금)부터 근무일 5일
    assert r["reasons"] == ["MATERIAL_LEADTIME", "DUE_DATE_MISSED"]


def test_order_from_stock_feasible(X):
    r = next(o for o in X["order_feasibility"] if o["order_id"] == "O001")
    assert d(r["production_qty"]) == 0 and r["feasible"] is True


def test_open_order_filter(X):
    ids = {o["order_id"] for o in X["order_feasibility"]}
    assert len(ids) == 18 and "O018" not in ids and "O019" not in ids and "O020" in ids   # 취소·완료 제외, 마지막 날 포함


def test_outsourcing_cost_hand(X):
    # 외주안 v2: FG003 주별 수량 80/250/210/200 중 외주처 능력 150까지만 외주 → 80+150+150+150 = 530, 단가 38,000원
    out = X["plan_evaluation"]["PLAN-OUT"]
    assert sum(d(v) for v in out["outsourced_qty"].values()) == 80 + 150 + 150 + 150
    assert d(out["costs_krw"]["outsourcing"]) == 530 * 38000
    assert X["plan_evaluation"]["PLAN-REG"]["costs_krw"]["inventory_holding"] == "미산정"


def test_overtime_plan_hand(X):
    ot = X["plan_evaluation"]["PLAN-OT"]
    cap = {(c["resource_id"], c["week_start"]): c for c in ot["capacity"]}
    # R2 12/07주 필요시간: FG001 180×0.2+0.5 + FG002 100×0.25+0.5 + FG003 200×0.15+0.5
    assert d(cap[("R2", "2026-12-07")]["required_h"]) == Decimal("36.5") + Decimal("25.5") + Decimal("30.5")   # 92.5
    # 잔업시간: R2 11/23주 86−80=6, R2 12/07주 min(12.5, 10)=10, R3 12/07주 min(18, 10)=10 → 26h
    assert d(ot["worked_overtime_h"]) == 6 + 10 + 10
    # 잔업비: R2 16h × 75,000 + R3 10h × 82,500
    assert d(ot["costs_krw"]["overtime"]) == 16 * 75000 + 10 * 82500
    # 능력 초과: R2 12/07주 92.5−90=2.5, R3 12/07주 58−50=8
    assert d(ot["capacity_excess_total_h"]) == Decimal("2.5") + 8


def test_regular_plan_late_order_hand(X):
    late = {x["order_id"]: d(x["shortfall_qty"]) for x in X["plan_evaluation"]["PLAN-REG"]["late_orders"]}
    # FG003: 11/24까지 공급 = 재고 80 + 1주차 생산 80(11/20 금 완료) = 160, 누적 수요 O003 120 + O017 100 = 220 → 60 부족
    assert late["O017"] == 220 - 160


def test_plan_alternatives_do_not_increase_excess(X):
    reg = {(c["resource_id"], c["week_start"]): d(c["excess_h"]) for c in X["plan_evaluation"]["PLAN-REG"]["capacity"]}
    for pid in ("PLAN-OT", "PLAN-OUT"):
        for c in X["plan_evaluation"][pid]["capacity"]:
            assert d(c["excess_h"]) <= reg[(c["resource_id"], c["week_start"])], (pid, c)
    assert d(X["plan_evaluation"]["PLAN-REG"]["worked_overtime_h"]) == 0


def test_v1_ground_truth_preserved(gt):
    prev = gt["dataset"]["previous_version"]
    assert gt["dataset"]["version"] == "v2" and prev["version"] == "v1"
    assert sha256_file(P.ground_truth_v1) == prev["sha256"]
    assert (P.results / "archive_v1" / "validation_report.json").exists()


@needs_uci
def test_uci_day_code_kept_and_dates_synthetic(gt):
    v = read_json(P.results / "uci_source_verification.json")
    assert "미확인" in v["day_sequence"]["code_meaning"]
    sales = gt["entities"]["SalesHistory"]
    assert all(s["date_is_synthetic"] is True and s["source_day_of_week_code"] in (2, 3, 4, 5, 6) for s in sales)


# ---------------------------------------------------------------------------
# E00 XML 보정
# ---------------------------------------------------------------------------
E00 = ROOT / "results" / "E00_environment"


def test_e00_xml_original_preserved_and_repaired_parses():
    import xml.etree.ElementTree as ET
    rec = read_json(E00 / "xml_repair_record.json")
    assert sha256_file(E00 / "pytest_report.xml") == rec["source"]["sha256"] == "e92fa8938bffedc56a26141c8e060c6bb1ff73ffeb5d8e27954d361fdf4560b3"
    assert rec["source"]["xml_parse_error"] is not None          # 원본은 깨진 상태 그대로 보존
    root = ET.parse(E00 / "pytest_report.repaired.xml").getroot()
    suite = root if root.tag == "testsuite" else root.find("testsuite")
    assert suite.get("tests") == "16" and suite.get("failures") == "0" and suite.get("errors") == "0"


def test_masked_xml_parses_without_host_or_home():
    # E01 pytest_report.xml은 이 테스트 실행 중에 만들어지므로 run_e01.py가 생성 직후 파싱 여부를 직접 확인한다.
    import socket
    import xml.etree.ElementTree as ET
    path = E00 / "pytest_report.repaired.xml"
    ET.parse(path)
    text = path.read_text(encoding="utf-8").lower()
    assert socket.gethostname().lower() not in text
    assert str(Path.home()).lower() not in text and str(Path.home()).replace("\\", "/").lower() not in text


def test_sanitize_junit_xml_keeps_xml_valid():
    import socket
    import xml.etree.ElementTree as ET
    import check_environment as ce
    home = str(Path.home())
    raw = (f'<?xml version="1.0" encoding="utf-8"?><testsuites><testsuite name="pytest" tests="1" hostname="{socket.gethostname()}">'
           f'<testcase classname="t" name="a" file="{home}\\x.py"><system-out>{home}\\log &amp; more</system-out></testcase>'
           f'</testsuite></testsuites>')
    out = ce.sanitize_junit_xml(raw)
    root = ET.fromstring(out)
    assert root.find("testsuite").get("hostname") == ce.MASKED_HOST
    assert root.find("testsuite/testcase").get("file") == "<HOME>\\x.py"     # 파싱 후 값은 원래 문자열
    assert home.lower() not in out.lower() and socket.gethostname().lower() not in out.lower()
    broken = raw.replace(f'hostname="{socket.gethostname()}"', 'hostname="<HOST>"').replace(home, "<HOME>")
    with pytest.raises(ET.ParseError):
        ET.fromstring(broken)
    ET.fromstring(ce.repair_masked_junit_xml(broken))


def test_cq10_long_leadtime(X):
    assert X["cq"]["CQ-10"]["products"] == {"FG003": ["M007"]}   # 우선 공급사 리드타임 21일 ≥ 14


# ---------------------------------------------------------------------------
# 데이터·검증 결과
# ---------------------------------------------------------------------------
def test_reference_integrity_clean(gt):
    E = gt["entities"]
    prods, mats = {p["product_id"] for p in E["Product"]}, {m["material_id"] for m in E["Material"]}
    res, boms = {r["resource_id"] for r in E["Resource"]}, {b["bom_id"] for b in E["BOM"]}
    assert all((l["component_material_id"] or l["component_product_id"]) in prods | mats and l["bom_id"] in boms for l in E["BOMLine"])
    assert all(o["product_id"] in prods for o in E["Order"]) and all(op["resource_id"] in res for op in E["Operation"])
    assert all(i["item_id"] in (prods if i["item_type"] == "PRODUCT" else mats) for i in E["Inventory"])


def test_injected_error_records_complete():
    inj = read_json(P.injected)
    for e in inj["errors"]:
        assert e["error_id"].startswith("E01-ERR-") and e["file"] and e["row_number"] >= 1 and e["field"] and e["error_code"].startswith("ERR-")
        assert e["expected_detection"]["rule"] and e["expected_detection"]["result"]
    assert len({e["error_id"] for e in inj["errors"]}) == 13


def test_record_counts(gt):
    c = gt["dataset"]["record_counts"]
    assert (c["Product"], c["Material"], c["Order"], c["Resource"], c["Supplier"], c["ProductionPlan"]) == (5, 10, 20, 3, 4, 3)
    assert 20 <= c["BOMLine"] <= 25


def test_validation_report_all_passed():
    rep = read_json(P.results / "validation_report.json")
    failed = [c for c in rep["checks"] if not c["passed"]]
    assert failed == [], failed
    assert rep["injection"]["injected"] == rep["injection"]["confirmed"] == 13
    assert set(rep["cq_status"].values()) == {"검증 통과"}


def test_verify_does_not_import_generation_modules():
    tree = ast.parse((ROOT / "scripts" / "e01_verify.py").read_text(encoding="utf-8"))
    names = {a.name for n in ast.walk(tree) if isinstance(n, ast.Import) for a in n.names}
    names |= {n.module for n in ast.walk(tree) if isinstance(n, ast.ImportFrom)}
    assert not names & {"e01_expected", "e01_generate", "e01_inject_errors"}


def test_vendor_encodings_and_leading_zeros():
    assert (P.vendor_a / "materials.csv").read_bytes().startswith(b"\xef\xbb\xbf")
    assert not (P.vendor_a / "orders.csv").read_bytes().startswith(b"\xef\xbb\xbf")
    with pytest.raises(UnicodeDecodeError):
        (P.vendor_b / "vendor_b_orders.csv").read_bytes().decode("utf-8")
    assert "주문번호" in (P.vendor_b / "vendor_b_orders.csv").read_bytes().decode("cp949")
    rows = list(csv.reader(io.StringIO((P.vendor_a / "products.csv").read_text(encoding="utf-8"))))
    assert rows[1][0] == "0001001"
    ws = openpyxl.load_workbook(P.vendor_a / "vendor_a_master.xlsx")["Products"]
    assert ws["A2"].value == "0001001"
    wb = openpyxl.load_workbook(P.vendor_b / "vendor_b_master.xlsx")
    assert len(wb.sheetnames) == 4 and wb["BOM"].merged_cells.ranges


def test_clean_vendor_files_untouched_by_injection():
    inj = read_json(P.injected)
    for e in inj["errors"]:
        assert e["file"].startswith("invalid/") and e["clean_file"] == e["file"][len("invalid/"):]


@needs_uci
def test_uci_raw_unmodified_and_untracked():
    rec = json.loads((P.uci_dir / "download_record.json").read_text(encoding="utf-8"))
    assert sha256_file(P.uci_dir / rec["zip"]["file"]) == rec["zip"]["sha256"]
    assert sha256_file(P.uci_csv) == rec["members"][0]["sha256"]
    r = subprocess.run(["git", "check-ignore", "-q", str(P.uci_csv)], cwd=ROOT)
    assert r.returncode == 0, "UCI 원본이 Git 추적 제외 대상이 아님"


@needs_uci
def test_uci_relationships_recorded():
    v = read_json(P.results / "uci_source_verification.json")
    assert v["row_count"] == 60 and v["column_count"] == 13 and v["delimiter"] == ";"
    rel = {c["relation"]: c["exact_match_rows"] for c in v["relationship_checks"]}
    assert rel["target_total == type_a + type_b + type_c"] == 60


# ---------------------------------------------------------------------------
# 음성 대조 (임시 복사본에서 실행)
# ---------------------------------------------------------------------------
@pytest.fixture()
def sandbox():
    with tempfile.TemporaryDirectory(prefix="e01_neg_") as t:
        root = Path(t)
        for sub in ("data/canonical", "data/vendor_a", "data/vendor_b", "data/invalid"):
            shutil.copytree(ROOT / sub, root / sub)
        stg = "data/raw/public/uci_daily_demand/staging_uci_daily_orders_v1.json"
        if (ROOT / stg).exists():
            (root / stg).parent.mkdir(parents=True)
            shutil.copy(ROOT / stg, root / stg)
        yield root


def _failed_sections(root: Path) -> set[str]:
    rep = e01_verify.run(Paths(root))
    return {c["section"] for c in rep["checks"] if not c["passed"]}


@needs_uci
def test_negative_vendor_value_change_detected(sandbox):
    f = sandbox / "data/vendor_a/orders.csv"
    f.write_text(f.read_text(encoding="utf-8").replace("SO260001,0001001,80,", "SO260001,0001001,81,"), encoding="utf-8", newline="")
    assert "vendor_files" in _failed_sections(sandbox)


@needs_uci
def test_negative_expected_value_change_detected(sandbox):
    target = Paths(sandbox).ground_truth  # 현재 정답셋 버전 (v2)
    g = read_json(target)
    g["expected"]["production_requirements"]["FG001"]["net_requirement"] = "461"
    write_json(target, g)
    assert "decision_calc" in _failed_sections(sandbox)


@needs_uci
def test_negative_injection_reverted_detected(sandbox):
    f = sandbox / "data/invalid/vendor_a/routing.csv"
    f.write_text(f.read_text(encoding="utf-8").replace(",R09,", ",R02,"), encoding="utf-8", newline="")
    assert "error_injection" in _failed_sections(sandbox)


@needs_uci
def test_negative_undeclared_change_detected(sandbox):
    f = sandbox / "data/invalid/vendor_a/suppliers.csv"
    f.write_text(f.read_text(encoding="utf-8").replace("SUP-01,", "SUP-01X,"), encoding="utf-8", newline="")
    assert "error_injection" in _failed_sections(sandbox)


@needs_uci
def test_regeneration_is_byte_identical():
    import run_e01
    with tempfile.TemporaryDirectory(prefix="e01_regen_") as t:
        run_e01.pipeline(Paths(Path(t)))
        rels = run_e01.GENERATED + run_e01.GENERATED_INVALID
        diff = [r for r in rels if sha256_file(Path(t) / r) != sha256_file(ROOT / r)]
    assert diff == []
