"""E03 문서 샘플 생성 (DOCX 2, 텍스트형 PDF 2, PPTX 2).

    .\\.venv\\Scripts\\python.exe scripts\\e03_generate_docs.py [--out-root DIR]

- 원천: data/canonical/e03_document_truth_v1.json (정답 명세, 수정하지 않음) — 기준 해시가 다르면 중단.
- 출력: data/documents/<문서>.docx|pdf|pptx, data/canonical/e03_document_locations_v1.json (생성 시 실제 위치).
- 수치·텍스트는 명세의 문자열을 그대로 쓴다. 새 값을 만들지 않는다 (PDF 페이지 번호 표기 제외 — 아래 PAGE_LABEL).
- 문서 속성(작성자 등)은 고정값으로 쓰고, OOXML zip 항목 시각과 PDF 생성 시각을 고정해 재현성을 확보한다.
- 검증은 e03_validate_docs.py가 저장된 파일을 다시 읽어 독립적으로 수행한다 (이 모듈을 import하지 않음).
"""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import os
import sys
import time
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt as DPt
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches
from pptx.util import Pt as PPt
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

from e01_common import normalize_xlsx

ROOT = Path(__file__).resolve().parents[1]
TRUTH_SHA256 = "982af3be7af0b4f512851a3a2cba2d68578b8e046540d94f0efe1170b0cf554a"
GT_SHA256 = "262dbe75dbfaf8c0c32003144d62e0b2057c61187dce4b548a457cf8819fc48e"
TARGET_DOCS = ["D-01", "D-02", "P-01", "P-02", "T-01", "T-02"]
FIXED_DT = dt.datetime(2026, 10, 20, 0, 0, 0)
AUTHOR = "chungpamo-poc (synthetic sample generator)"
FONT_NAME = "Malgun Gothic"            # DOCX·PPTX: 글꼴 이름만 지정 (임베드하지 않음)
PDF_FONT_FILE = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts" / "malgun.ttf"  # PDF: 서브셋 임베드 (fsType 8)
PAGE_LABEL = "{n} / {total}"           # PDF 페이지 번호 표기 (명세 외 표시 요소)
# 텍스트 블록은 11개(CS1 4 + CS2 4 + CS3 3). E03 1단계 보고의 '13개'는 오기였다 (정답 명세 파일 기준으로 정정).
EXPECTED_COUNTS = {"content_sets": 3, "tables": 6, "cells": 361, "values": 142, "blocks": 11}


# ---------------------------------------------------------------------------
# 입력 검증
# ---------------------------------------------------------------------------
def load_and_verify() -> dict:
    gt_path = ROOT / "data" / "canonical" / "ground_truth_v2.json"
    tr_path = ROOT / "data" / "canonical" / "e03_document_truth_v1.json"
    for p, want in ((gt_path, GT_SHA256), (tr_path, TRUTH_SHA256)):
        got = hashlib.sha256(p.read_bytes()).hexdigest()
        if got != want:
            raise SystemExit(f"기준 해시 불일치 → 생성 중단: {p.name} {got}")
    gt = json.loads(gt_path.read_text(encoding="utf-8"))
    truth = json.loads(tr_path.read_text(encoding="utf-8"))
    sets = truth["content_sets"]
    counts = {"content_sets": len(sets), "tables": sum(len(s["tables"]) for s in sets.values()),
              "cells": sum(len(t["cells"]) for s in sets.values() for t in s["tables"]),
              "values": sum(len(s["extracted_values"]) for s in sets.values()),
              "blocks": sum(len(s["blocks"]) for s in sets.values())}
    if counts != EXPECTED_COUNTS:
        raise SystemExit(f"정답 명세 개수 불일치 → 중단: {counts}")
    uids = {r["record_uid"] for rows in gt["entities"].values() for r in rows}
    missing = []
    for s in sets.values():
        for v in s["extracted_values"]:
            ref = v["canonical_ref"]
            if "record_uid" in ref and ref["record_uid"] not in uids:
                missing.append(ref)
            if "expected_path" in ref:
                node = gt
                try:
                    for k in ref["expected_path"].split("."):
                        node = node[k]
                except KeyError:
                    missing.append(ref)
    if missing:
        raise SystemExit(f"canonical 참조 누락 → 중단: {missing[:3]}")
    return truth


def items_in_order(cs: dict, order: list[str]) -> list[tuple[str, dict]]:
    blocks = {b["block_id"]: b for b in cs["blocks"]}
    tables = {t["table_id"]: t for t in cs["tables"]}
    return [("block", blocks[i]) if i in blocks else ("table", tables[i]) for i in order]


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------
def _docx_font(doc) -> None:
    st = doc.styles["Normal"]
    st.font.name = FONT_NAME
    st.font.size = DPt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)


def _docx_page_field(paragraph) -> None:
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    paragraph._p.append(fld)


def _body_index(doc) -> int:
    return sum(1 for el in doc.element.body if el.tag in (qn("w:p"), qn("w:tbl"))) - 1


def make_docx(spec: dict, cs: dict, path: Path) -> dict:
    doc = Document()
    _docx_font(doc)
    cp = doc.core_properties
    cp.author = cp.last_modified_by = AUTHOR
    cp.title = cs["title"]
    cp.subject = f"E03 synthetic sample {spec['doc_id']} ({cs['content_set']})"
    cp.comments = "Synthetic data generated from E01 canonical v2. Not real company records."
    cp.keywords = cp.category = ""
    cp.created = cp.modified = FIXED_DT
    cp.last_printed = FIXED_DT
    cp.revision = 1
    loc = {"blocks": {}, "tables": {}}
    layout = spec["layout"]
    blocks = {b["block_id"]: b for b in cs["blocks"]}
    sec = doc.sections[0]
    if "header" in layout and layout["header"].split()[0] in blocks:
        bid = layout["header"].split()[0]
        p = sec.header.paragraphs[0]
        p.text = blocks[bid]["text"]
        loc["blocks"][bid] = [{"format": "docx", "part": "header", "section_index": 0, "paragraph_index": 0}]
    fp = sec.footer.paragraphs[0]
    footer_bid = layout.get("footer", "").split()[0] if layout.get("footer", "").split() else ""
    if footer_bid in blocks:
        fp.text = blocks[footer_bid]["text"]
        loc["blocks"][footer_bid] = [{"format": "docx", "part": "footer", "section_index": 0, "paragraph_index": 0}]
        fp = sec.footer.add_paragraph()
    _docx_page_field(fp)
    loc["page_number_field"] = {"part": "footer", "section_index": 0}
    table_index = 0
    for kind, obj in items_in_order(cs, layout["body_order"]):
        if kind == "block":
            if obj["type"] == "heading":
                doc.add_heading(obj["text"], level=1)
            else:
                doc.add_paragraph(obj["text"])
            loc["blocks"].setdefault(obj["block_id"], []).append(
                {"format": "docx", "part": "body", "body_index": _body_index(doc), "note": "각주 블록은 본문 문단으로 렌더링 (Word 각주 part 미사용)" if obj["type"] == "footnote" else None})
        else:
            doc.add_paragraph(obj["caption"])
            cap_idx = _body_index(doc)
            t = doc.add_table(rows=obj["n_rows"], cols=obj["n_cols"])
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for m in obj["merged_ranges"]:
                a = t.cell(m["row"] - 1, m["col"] - 1)
                b = t.cell(m["row"] - 1 + m["rowspan"] - 1, m["col"] - 1 + m["colspan"] - 1)
                a.merge(b)
            for c in obj["cells"]:
                cell = t.cell(c["row"] - 1, c["col"] - 1)
                cell.text = c["text"]
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = DPt(8.5)
                        run.font.bold = c["is_header"]
            loc["tables"][obj["table_id"]] = {"fragments": [{
                "fragment_index": 0, "format": "docx", "part": "body", "body_index": _body_index(doc), "table_index": table_index,
                "caption_location": {"part": "body", "body_index": cap_idx},
                "logical_rows": list(range(1, obj["n_rows"] + 1)),
                "cells": {f"{c['row']},{c['col']}": {"row": c["row"], "col": c["col"]} for c in obj["cells"]}}]}
            table_index += 1
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    normalize_xlsx(path)  # OOXML 공통: zip 항목 시각·문서 속성 시각 고정
    return loc


# ---------------------------------------------------------------------------
# PDF (reportlab canvas, 좌표 직접 계산)
# ---------------------------------------------------------------------------
PW, PH = A4
MARGIN = 50.0
AVAIL = PW - 2 * MARGIN
PDF_FONT = "MalgunGothic"


def wrap(text: str, size: float, width: float) -> list[str]:
    """줄을 이어 붙이면 원문과 같아지도록 자른다 (공백은 앞줄 끝에 남김)."""
    lines, start = [], 0
    while start < len(text):
        end = start
        last_space = -1
        while end < len(text) and pdfmetrics.stringWidth(text[start:end + 1], PDF_FONT, size) <= width:
            if text[end] == " ":
                last_space = end
            end += 1
        if end < len(text) and last_space > start:
            end = last_space + 1
        if end == start:
            end = start + 1
        lines.append(text[start:end])
        start = end
    return lines


class PdfWriter:
    def __init__(self, path: Path, title: str, subject: str, total_pages: int, header_text=None, footer_text=None) -> None:
        self.c = canvas.Canvas(str(path), pagesize=A4, invariant=1)
        self.c.setTitle(title)
        self.c.setAuthor(AUTHOR)
        self.c.setSubject(subject)
        self.c.setCreator("scripts/e03_generate_docs.py")
        self.c.setKeywords("synthetic, E03")
        self.page = 1
        self.total = total_pages
        self.y = PH - MARGIN       # 다음 출력 위치 (top-left 원점 환산 전 reportlab 좌표)
        self.header_text, self.footer_text = header_text, footer_text
        self.header_locs, self.footer_locs = [], []
        self._page_furniture()

    def bbox(self, x0, y_base, x1, size) -> list[float]:
        """reportlab 좌표 → top-left 원점 bbox [x0, top, x1, bottom] (pt)."""
        top = PH - (y_base + size * 0.88)
        bottom = PH - (y_base - size * 0.22)
        return [round(x0, 2), round(top, 2), round(x1, 2), round(bottom, 2)]

    def text_line(self, s: str, x: float, y_base: float, size: float) -> list[float]:
        self.c.setFont(PDF_FONT, size)
        self.c.drawString(x, y_base, s)
        return self.bbox(x, y_base, x + pdfmetrics.stringWidth(s, PDF_FONT, size), size)

    def _page_furniture(self) -> None:
        if self.header_text:
            b = self.text_line(self.header_text, MARGIN, PH - 30, 8)
            self.header_locs.append({"format": "pdf", "page": self.page, "bbox": b})
        if self.footer_text:
            b = self.text_line(self.footer_text, MARGIN, 34, 7.5)
            self.footer_locs.append({"format": "pdf", "page": self.page, "bbox": b})
        self.c.setFont(PDF_FONT, 8)
        self.c.drawCentredString(PW / 2, 20, PAGE_LABEL.format(n=self.page, total=self.total))
        self.y = PH - MARGIN - (12 if self.header_text else 0)

    def new_page(self) -> None:
        self.c.showPage()
        self.page += 1
        self._page_furniture()

    def paragraph(self, text: str, size: float) -> dict:
        lines = wrap(text, size, AVAIL)
        lead = size * 1.45
        boxes = []
        for ln in lines:
            self.y -= lead
            boxes.append(self.text_line(ln, MARGIN, self.y, size))
        self.y -= size * 0.6
        box = [min(b[0] for b in boxes), min(b[1] for b in boxes), max(b[2] for b in boxes), max(b[3] for b in boxes)]
        return {"format": "pdf", "page": self.page, "bbox": box, "lines": len(lines)}

    def table_fragment(self, t: dict, rows: list[int], size: float, widths: list[float], fragment_index: int) -> dict:
        cells = [c for c in t["cells"] if c["row"] in rows]
        row_h = size * 2.0
        x_off = [MARGIN]
        for w in widths:
            x_off.append(x_off[-1] + w)
        y_top = self.y - 4          # reportlab 좌표의 표 윗변
        row_top = {r: y_top - i * row_h for i, r in enumerate(rows)}
        cell_locs = {}
        self.c.setLineWidth(0.5)
        for c in cells:
            x0, x1 = x_off[c["col"] - 1], x_off[c["col"] - 1 + c["colspan"]]
            top = row_top[c["row"]]
            bottom = top - row_h * c["rowspan"]
            self.c.rect(x0, bottom, x1 - x0, top - bottom, stroke=1, fill=0)
            base = bottom + (top - bottom - size) / 2 + size * 0.22
            if c["text"]:
                self.c.setFont(PDF_FONT, size)
                self.c.drawString(x0 + 3, base, c["text"])
            cell_locs[f"{c['row']},{c['col']}"] = {"page": self.page,
                                                   "bbox": [round(x0, 2), round(PH - top, 2), round(x1, 2), round(PH - bottom, 2)]}
        bottom_all = y_top - row_h * len(rows)
        self.y = bottom_all - 8
        return {"fragment_index": fragment_index, "format": "pdf", "page": self.page,
                "bbox": [round(MARGIN, 2), round(PH - y_top, 2), round(x_off[-1], 2), round(PH - bottom_all, 2)],
                "logical_rows": rows, "repeated_header_rows": [r for r in rows if r <= t["header_rows"]] if fragment_index > 0 else [],
                "cells": cell_locs}

    def save(self) -> None:
        self.c.save()


def column_widths(t: dict) -> tuple[list[float], float]:
    """모든 행 기준으로 열 너비를 정하고, 가용 폭을 넘으면 글자 크기를 줄인다."""
    size = 8.5
    while True:
        w = [0.0] * t["n_cols"]
        for c in t["cells"]:
            if c["colspan"] == 1:
                w[c["col"] - 1] = max(w[c["col"] - 1], pdfmetrics.stringWidth(c["text"], PDF_FONT, size) + 8)
        for c in t["cells"]:
            if c["colspan"] > 1:
                need = pdfmetrics.stringWidth(c["text"], PDF_FONT, size) + 8
                span = range(c["col"] - 1, c["col"] - 1 + c["colspan"])
                have = sum(w[i] for i in span)
                if have < need:
                    for i in span:
                        w[i] += (need - have) / c["colspan"]
        total = sum(w)
        if total <= AVAIL or size <= 6:
            extra = (AVAIL - total) / len(w) if total < AVAIL else 0
            return [x + extra for x in w], size
        size -= 0.25


def make_pdf(spec: dict, cs: dict, path: Path) -> dict:
    blocks = {b["block_id"]: b for b in cs["blocks"]}
    tables = {t["table_id"]: t for t in cs["tables"]}
    loc = {"blocks": {}, "tables": {}, "coordinate_system": {"unit": "pt", "origin": "top-left", "page_size": [round(PW, 2), round(PH, 2)],
                                                            "bbox": "[x0, top, x1, bottom]"}}
    path.parent.mkdir(parents=True, exist_ok=True)
    if spec["doc_id"] == "P-01":
        w = PdfWriter(path, cs["title"], "E03 synthetic sample P-01 (CS1)", 1, header_text=blocks["CS1-b02"]["text"])
        plan = ["CS1-b01", "CS1-b03", "CS1-T1", "CS1-T2", "CS1-b04"]
        splits = {}
    else:
        w = PdfWriter(path, cs["title"], "E03 synthetic sample P-02 (CS2)", 3, footer_text=blocks["CS2-b04"]["text"])
        t1 = tables["CS2-T1"]
        data = list(range(t1["header_rows"] + 1, t1["n_rows"] + 1))
        splits = {"CS2-T1": [data[:12], data[12:]]}      # 명세 split_policy: 데이터 12행마다 넘김, 헤더 반복
        plan = ["CS2-b01", "CS2-b02", ("CS2-T1", 0), "PAGE", ("CS2-T1", 1), "CS2-b03", "PAGE", "CS2-T2"]
    for item in plan:
        if item == "PAGE":
            w.new_page()
            continue
        tid, frag = (item if isinstance(item, tuple) else (item, None))
        if tid in blocks:
            b = blocks[tid]
            size = 13 if b["type"] == "heading" else 9
            loc["blocks"].setdefault(tid, []).append(w.paragraph(b["text"], size))
            continue
        t = tables[tid]
        widths, size = column_widths(t)
        hdr = list(range(1, t["header_rows"] + 1))
        if frag is None or frag == 0:
            cap = w.paragraph(t["caption"], 9)
        rows = hdr + (splits[tid][frag] if tid in splits else list(range(t["header_rows"] + 1, t["n_rows"] + 1)))
        f = w.table_fragment(t, rows, size, widths, frag or 0)
        entry = loc["tables"].setdefault(tid, {"fragments": [], "caption_location": cap if (frag is None or frag == 0) else None,
                                                "font_size": size})
        entry["fragments"].append(f)
    if w.header_locs:
        loc["blocks"]["CS1-b02"] = w.header_locs
    if w.footer_locs:
        loc["blocks"]["CS2-b04"] = w.footer_locs
    loc["pages"] = w.page
    w.save()
    return loc


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------
def _pptx_fonts(text_frame, size: float, bold=False) -> None:
    for p in text_frame.paragraphs:
        for r in p.runs:
            r.font.size = PPt(size)
            r.font.bold = bold
            r.font.name = FONT_NAME
            rpr = r._r.get_or_add_rPr()
            ea = rpr.find(qn("a:ea"))
            if ea is None:
                ea = OxmlElement("a:ea")
                rpr.append(ea)
            ea.set("typeface", FONT_NAME)


def _textbox(shapes, name: str, text: str, left, top, width, height, size: float, bold=False):
    tb = shapes.add_textbox(left, top, width, height)
    tb.name = name
    tb.text_frame.word_wrap = True
    tb.text_frame.text = text
    _pptx_fonts(tb.text_frame, size, bold)
    return tb


def _table(shapes, name: str, t: dict, rows: list[int], left, top, width):
    gf = shapes.add_table(len(rows), t["n_cols"], left, top, width, Inches(0.4) * len(rows))
    gf.name = name
    tbl = gf.table
    idx = {r: i for i, r in enumerate(rows)}
    for m in t["merged_ranges"]:
        if m["row"] in idx:
            tbl.cell(idx[m["row"]], m["col"] - 1).merge(tbl.cell(idx[m["row"]] + m["rowspan"] - 1, m["col"] - 1 + m["colspan"] - 1))
    cell_locs = {}
    for c in t["cells"]:
        if c["row"] not in idx:
            continue
        cell = tbl.cell(idx[c["row"]], c["col"] - 1)
        cell.text = c["text"]
        _pptx_fonts(cell.text_frame, 11, c["is_header"])
        cell_locs[f"{c['row']},{c['col']}"] = {"table_row": idx[c["row"]] + 1, "table_col": c["col"]}
    return gf, cell_locs


def make_pptx(spec: dict, cs: dict, path: Path) -> dict:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    cp = prs.core_properties
    cp.author = cp.last_modified_by = AUTHOR
    cp.title = cs["title"]
    cp.subject = f"E03 synthetic sample {spec['doc_id']} ({cs['content_set']})"
    cp.comments = "Synthetic data generated from E01 canonical v2. Not real company records."
    cp.keywords = cp.category = ""
    cp.created = cp.modified = cp.last_printed = FIXED_DT
    cp.revision = 1
    blank = prs.slide_layouts[6]
    blocks = {b["block_id"]: b for b in cs["blocks"]}
    tables = {t["table_id"]: t for t in cs["tables"]}
    loc = {"blocks": {}, "tables": {}, "slides": 0,
           "coordinate_system": {"unit": "EMU", "note": "1 inch = 914400 EMU, 위치는 slide 좌상단 기준 left/top/width/height"}}
    L, W = Inches(0.6), Inches(12.1)

    def shape_loc(slide_no, shp, nested=None, part="slide"):
        return {"format": "pptx", "part": part, "slide_index": slide_no, "shape_id": shp.shape_id, "shape_name": shp.name,
                "nested_path": nested or [shp.name], "left": int(shp.left), "top": int(shp.top), "width": int(shp.width), "height": int(shp.height)}

    def add_table_frag(slide, slide_no, tid, rows, frag_idx, top):
        t = tables[tid]
        cap = _textbox(slide.shapes, f"cap_{tid}_f{frag_idx}", t["caption"], L, top, W, Inches(0.4), 12, True) if frag_idx == 0 else None
        gf, cells = _table(slide.shapes, f"tbl_{tid}_f{frag_idx}", t, rows, L, top + Inches(0.5), W)
        e = loc["tables"].setdefault(tid, {"fragments": []})
        e["fragments"].append({"fragment_index": frag_idx, **shape_loc(slide_no, gf), "logical_rows": rows,
                               "repeated_header_rows": [r for r in rows if r <= t["header_rows"]] if frag_idx > 0 else [],
                               "caption_location": shape_loc(slide_no, cap) if cap else None, "cells": cells})

    if spec["doc_id"] == "T-01":
        t1 = tables["CS3-T1"]
        hdr = list(range(1, t1["header_rows"] + 1))
        fg = [r for r in range(t1["header_rows"] + 1, t1["n_rows"] + 1) if next(c for c in t1["cells"] if c["row"] == r and c["col"] == 1)["text"].startswith("FG")]
        sf = [r for r in range(t1["header_rows"] + 1, t1["n_rows"] + 1) if r not in fg]
        s1 = prs.slides.add_slide(blank)
        tb = _textbox(s1.shapes, "title_CS3-b01", blocks["CS3-b01"]["text"], L, Inches(3.0), W, Inches(1.2), 28, True)
        loc["blocks"]["CS3-b01"] = [shape_loc(1, tb)]
        s2 = prs.slides.add_slide(blank)
        add_table_frag(s2, 2, "CS3-T1", hdr + fg, 0, Inches(0.6))
        s3 = prs.slides.add_slide(blank)
        add_table_frag(s3, 3, "CS3-T1", hdr + sf, 1, Inches(0.6))
        s4 = prs.slides.add_slide(blank)
        add_table_frag(s4, 4, "CS3-T2", list(range(1, tables["CS3-T2"]["n_rows"] + 1)), 0, Inches(0.6))
        grp = s4.shapes.add_group_shape()
        grp.name = "grp_CS3-b02"
        box = grp.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, L, Inches(5.6), W, Inches(1.1))
        box.name = "bg_CS3-b02"
        tb = _textbox(grp.shapes, "txt_CS3-b02", blocks["CS3-b02"]["text"], L + Inches(0.2), Inches(5.7), W - Inches(0.4), Inches(0.9), 14)
        loc["blocks"]["CS3-b02"] = [shape_loc(4, tb, nested=[grp.name, tb.name])]
        loc["groups"] = [{"slide_index": 4, "group_shape_id": grp.shape_id, "group_name": grp.name,
                          "children": [{"shape_id": box.shape_id, "name": box.name}, {"shape_id": tb.shape_id, "name": tb.name}]}]
        s4.notes_slide.notes_text_frame.text = blocks["CS3-b03"]["text"]
        loc["blocks"]["CS3-b03"] = [{"format": "pptx", "part": "notes", "slide_index": 4, "nested_path": ["notes"]}]
        loc["slides"] = 4
    else:  # T-02: CS1 (명세 layout: 제목·문서번호 / CS1-T1 / CS1-T2 + CS1-b03)
        s1 = prs.slides.add_slide(blank)
        tb = _textbox(s1.shapes, "title_CS1-b01", blocks["CS1-b01"]["text"], L, Inches(2.6), W, Inches(1.2), 26, True)
        loc["blocks"]["CS1-b01"] = [shape_loc(1, tb)]
        tb = _textbox(s1.shapes, "txt_CS1-b02", blocks["CS1-b02"]["text"], L, Inches(4.0), W, Inches(0.6), 14)
        loc["blocks"]["CS1-b02"] = [shape_loc(1, tb)]
        s2 = prs.slides.add_slide(blank)
        add_table_frag(s2, 2, "CS1-T1", list(range(1, tables["CS1-T1"]["n_rows"] + 1)), 0, Inches(0.6))
        s3 = prs.slides.add_slide(blank)
        add_table_frag(s3, 3, "CS1-T2", list(range(1, tables["CS1-T2"]["n_rows"] + 1)), 0, Inches(0.6))
        tb = _textbox(s3.shapes, "txt_CS1-b03", blocks["CS1-b03"]["text"], L, Inches(4.6), W, Inches(1.4), 13)
        loc["blocks"]["CS1-b03"] = [shape_loc(3, tb)]
        loc["slides"] = 3
        loc["not_rendered"] = {"CS1-b04": "명세 layout(T-02)에 포함되지 않아 렌더링하지 않음"}
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    normalize_xlsx(path)
    return loc


# ---------------------------------------------------------------------------
def generate(out_root: Path | None = None) -> dict:
    truth = load_and_verify()
    root = Path(out_root) if out_root else ROOT
    if not PDF_FONT_FILE.exists():
        raise SystemExit(f"PDF 한글 폰트 없음: {PDF_FONT_FILE.name}")
    if PDF_FONT not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(PDF_FONT, str(PDF_FONT_FILE)))
    out = {"schema": "chungpamo.e03_document_locations", "version": "v1",
           "truth": {"file": "data/canonical/e03_document_truth_v1.json", "sha256": TRUTH_SHA256},
           "note": "생성기가 렌더링하면서 기록한 실제 위치 (actual_location). 파서 추출 위치가 아니다.",
           "pdf_font": {"name": "Malgun Gothic (malgun.ttf, 로컬 시스템 폰트)", "os2_fsType": 8, "embedding": "서브셋 임베드",
                        "font_file_in_repo": False},
           "documents": {}}
    perf = {}
    for d in truth["documents"]:
        if d["doc_id"] not in TARGET_DOCS:
            continue
        cs = truth["content_sets"][d["content_set"]]
        path = root / d["file"]
        t0 = time.perf_counter()
        if d["format"] == "docx":
            loc = make_docx(d, cs, path)
        elif d["format"] == "pdf_text":
            loc = make_pdf(d, cs, path)
        else:
            loc = make_pptx(d, cs, path)
        perf[d["doc_id"]] = time.perf_counter() - t0
        out["documents"][d["doc_id"]] = {"file": d["file"], "format": d["format"], "content_set": d["content_set"], **loc}
    lp = root / "data" / "canonical" / "e03_document_locations_v1.json"
    lp.parent.mkdir(parents=True, exist_ok=True)
    lp.write_text(json.dumps(out, ensure_ascii=False, indent=1) + "\n", encoding="utf-8", newline="\n")
    return {"generation_seconds": perf}


def main(argv=None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out-root", type=Path, default=None)
    args = ap.parse_args(argv)
    print(generate(args.out_root))
    return 0


if __name__ == "__main__":
    sys.exit(main())
