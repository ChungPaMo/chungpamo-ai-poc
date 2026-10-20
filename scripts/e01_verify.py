"""E01 독립 검증 — 구현 B.

    .\\.venv\\Scripts\\python.exe scripts\\e01_verify.py [--out-root DIR]

e01_generate / e01_expected / e01_inject_errors를 import하지 않는다. 공유하는 것은 입출력 도구(e01_common)뿐이다.
검사 항목:
  1. canonical 무결성 (키·참조·단위·기간·BOM 구조·재고 수량·UCI 배분 합계)
  2. Vendor A/B 파일을 자체 파서로 읽어 canonical과 업무 의미가 같은지 대조
  3. 오류 주입: 정상본↔주입본 차이가 주입 목록과 정확히 같은지 + 오류 유형별 판정 조건 성립 여부
  4. 기대값 재계산 (위상 정렬 기반 다단계 전개 등 구현 A와 다른 방식) 후 ground_truth.expected와 대조
결과: results/E01_dataset/validation_report.json
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import io
import json
import sys
from collections import Counter, defaultdict, deque
import math
from decimal import Decimal, InvalidOperation
from fractions import Fraction
from pathlib import Path

import openpyxl

from e01_common import Paths, fmt, read_json, write_json

Z = Decimal(0)


def dec(x) -> Decimal:
    return Decimal(str(x))


def day(x: str) -> dt.date:
    return dt.date.fromisoformat(x)


class Checks:
    def __init__(self) -> None:
        self.items: list[dict] = []

    def add(self, section: str, name: str, ok: bool, detail=None) -> bool:
        self.items.append({"section": section, "check": name, "passed": bool(ok), "detail": detail})
        return ok

    def summary(self) -> dict:
        by = defaultdict(lambda: {"passed": 0, "failed": 0})
        for i in self.items:
            by[i["section"]]["passed" if i["passed"] else "failed"] += 1
        return {"total": len(self.items), "passed": sum(i["passed"] for i in self.items),
                "failed": sum(not i["passed"] for i in self.items), "by_section": dict(by)}


# ===========================================================================
# 1. canonical 무결성
# ===========================================================================
PK = {"Unit": ("unit_code",), "Product": ("product_id",), "Material": ("material_id",), "Supplier": ("supplier_id",),
      "SupplierMaterial": ("supplier_id", "material_id"), "BOM": ("bom_id",), "BOMLine": ("bom_id", "line_no"),
      "Order": ("order_id",), "SalesHistory": ("product_id", "period_start", "period_grain"),
      "Inventory": ("item_type", "item_id", "as_of_date"), "Resource": ("resource_id",), "Operation": ("operation_id",),
      "ProductionPlan": ("plan_id",), "ProductionPlanLine": ("plan_id", "line_no")}


def check_canonical(gt: dict, paths: Paths, C: Checks) -> None:
    E, P = gt["entities"], gt["parameters"]
    S = "canonical"
    C.add(S, "엔터티 14종 존재", sorted(E) == sorted(PK), sorted(E))
    for ent, keys in PK.items():
        ks = [tuple(r[k] for k in keys) for r in E[ent]]
        dup = [k for k, n in Counter(ks).items() if n > 1]
        C.add(S, f"{ent} 기본키 유일", not dup, dup)
    uids = [r["record_uid"] for rows in E.values() for r in rows]
    C.add(S, "record_uid 전체 유일", len(uids) == len(set(uids)), len(uids))
    ids = {ent: {r[PK[ent][0]] for r in E[ent]} for ent in ("Unit", "Product", "Material", "Supplier", "BOM", "Resource", "ProductionPlan")}
    unit = {u["unit_code"]: u for u in E["Unit"]}
    dim = lambda u: unit[u]["dimension"]
    base_of = {**{p["product_id"]: p["base_unit"] for p in E["Product"]}, **{m["material_id"]: m["base_unit"] for m in E["Material"]}}
    bad = []
    for u in E["Unit"]:
        if u["base_unit_code"] not in ids["Unit"] or dim(u["base_unit_code"]) != u["dimension"]:
            bad.append(u["unit_code"])
    C.add(S, "Unit 기준 단위 차원 일치", not bad, bad)
    refs = []
    for sm in E["SupplierMaterial"]:
        refs += [("SupplierMaterial", sm["supplier_id"] in ids["Supplier"]), ("SupplierMaterial", sm["material_id"] in ids["Material"])]
    for b in E["BOM"]:
        refs.append(("BOM", b["product_id"] in ids["Product"]))
    for ln in E["BOMLine"]:
        comp = ln["component_material_id"] or ln["component_product_id"]
        refs += [("BOMLine", ln["bom_id"] in ids["BOM"]), ("BOMLine", comp in (ids["Material"] | ids["Product"]))]
    for o in E["Order"]:
        refs.append(("Order", o["product_id"] in ids["Product"]))
    for i in E["Inventory"]:
        refs.append(("Inventory", i["item_id"] in (ids["Product"] if i["item_type"] == "PRODUCT" else ids["Material"])))
    for op in E["Operation"]:
        refs += [("Operation", op["product_id"] in ids["Product"]), ("Operation", op["resource_id"] in ids["Resource"])]
    for s in E["SalesHistory"]:
        refs.append(("SalesHistory", s["product_id"] in ids["Product"]))
    for l in E["ProductionPlanLine"]:
        refs += [("ProductionPlanLine", l["plan_id"] in ids["ProductionPlan"]), ("ProductionPlanLine", l["product_id"] in ids["Product"])]
    C.add(S, "참조 무결성 (FK)", all(ok for _, ok in refs), [e for e, ok in refs if not ok])
    bad = []
    for ln in E["BOMLine"]:
        comp = ln["component_material_id"] or ln["component_product_id"]
        if ln["unit"] not in unit or dim(ln["unit"]) != dim(base_of[comp]):
            bad.append((ln["bom_id"], ln["line_no"]))
    for i in E["Inventory"]:
        if dim(i["unit"]) != dim(base_of[i["item_id"]]):
            bad.append(i["item_id"])
    for o in E["Order"]:
        if dim(o["unit"]) != dim(base_of[o["product_id"]]):
            bad.append(o["order_id"])
    C.add(S, "수량 단위 차원 = 품목 기준 단위 차원", not bad, bad)
    bad = [o["order_id"] for o in E["Order"] if day(o["due_date"]) < day(o["order_date"])]
    bad += [b["bom_id"] for b in E["BOM"] if b["valid_to"] and day(b["valid_to"]) < day(b["valid_from"])]
    bad += [i["item_id"] for i in E["Inventory"] if day(i["as_of_date"]) > day(P["reference_date"])]
    bad += [s["product_id"] + s["period_start"] for s in E["SalesHistory"] if day(s["period_start"]) >= day(P["reference_date"])]
    C.add(S, "날짜 관계 (납기≥주문일, 유효기간, 기준일)", not bad, bad[:10])
    bad = [i["item_id"] for i in E["Inventory"] if dec(i["on_hand_qty"]) < 0 or dec(i["allocated_qty"]) > dec(i["on_hand_qty"])
           or (dec(i["on_order_qty"]) > 0 and not i["expected_receipt_date"])
           or (i["unit"] == "EA" and dec(i["on_hand_qty"]) != dec(i["on_hand_qty"]).to_integral_value())]
    C.add(S, "재고 수량 규칙", not bad, bad)
    # BOM 버전 유효 조건 중복 (정적 검사)
    act = [b for b in E["BOM"] if b["status"] == "APPROVED" and b["is_active"]]
    inf = Decimal("1e30")
    far = dt.date(9999, 12, 31)
    overlaps = []
    for i, x in enumerate(act):
        for y in act[i + 1:]:
            if x["product_id"] != y["product_id"]:
                continue
            period = day(x["valid_from"]) <= (day(y["valid_to"]) if y["valid_to"] else far) and \
                day(y["valid_from"]) <= (day(x["valid_to"]) if x["valid_to"] else far)
            qty = dec(x["min_qty"]) < (dec(y["max_qty"]) if y["max_qty"] else inf) and \
                dec(y["min_qty"]) < (dec(x["max_qty"]) if x["max_qty"] else inf)
            site = x["site_id"] is None or y["site_id"] is None or x["site_id"] == y["site_id"]
            if period and qty and site:
                overlaps.append((x["bom_id"], y["bom_id"]))
    C.add(S, "활성 BOM 버전 유효 조건 비중복 (QR-06 ①)", not overlaps, overlaps)
    C.add(S, "BOM 순환 없음 (QR-05)", not find_cycles(component_graph(E)), find_cycles(component_graph(E)))
    depth = max_depth(E)
    C.add(S, "BOM 최대 깊이 2 (ASM-04)", depth <= 2, depth)
    routing = defaultdict(set)
    for op in E["Operation"]:
        routing[op["product_id"]].add(op["resource_id"])
    bad = [l["line_no"] for l in E["ProductionPlanLine"] if l["resource_id"] and l["resource_id"] not in routing[l["product_id"]]]
    bad += [p for p in ids["Product"] if not routing[p]]
    C.add(S, "제품–설비 관계 (계획 자원 ∈ 공정 자원, 모든 제품 공정 보유)", not bad, bad)
    counts = {k: len(v) for k, v in E.items()}
    scale = {"Product": 5, "Material": 10, "Order": 20, "Resource": 3, "Supplier": 4, "ProductionPlan": 3}
    C.add(S, "초기 규모 (제품5·자재10·주문20·자원3·공급사4·계획안3)", all(counts[k] == v for k, v in scale.items()), counts)
    C.add(S, "BOM 라인 약 20행 (20~25)", 20 <= counts["BOMLine"] <= 25, counts["BOMLine"])
    fg = [p for p in E["Product"] if p["product_type"] == "FINISHED"]
    C.add(S, "완제품 3·반제품 2", len(fg) == 3 and counts["Product"] - len(fg) == 2, len(fg))
    weeks = sorted({l["period_start"] for l in E["ProductionPlanLine"]})
    C.add(S, "계획 4주 (주 시작일 4개, 월요일)", len(weeks) == 4 and all(day(w).weekday() == 0 for w in weeks), weeks)
    check_uci_allocation(E, paths, C)


def component_graph(E: dict) -> dict[str, set[str]]:
    by_bom = {b["bom_id"]: b["product_id"] for b in E["BOM"]}
    g = defaultdict(set)
    for ln in E["BOMLine"]:
        if ln["component_product_id"]:
            g[by_bom[ln["bom_id"]]].add(ln["component_product_id"])
    return g


def find_cycles(g: dict[str, set[str]]) -> list[list[str]]:
    cycles, color = [], {}

    def visit(n, path):
        color[n] = 1
        for c in sorted(g.get(n, ())):
            if color.get(c) == 1:
                cycles.append(path[path.index(c):] + [c])
            elif color.get(c) is None:
                visit(c, path + [c])
        color[n] = 2

    for n in sorted(g):
        if color.get(n) is None:
            visit(n, [n])
    return cycles


def max_depth(E: dict) -> int:
    g = component_graph(E)
    mat_parents = {b["product_id"] for b in E["BOM"]}

    def depth(n, seen=()):
        kids = g.get(n, set())
        own = 1 if n in mat_parents else 0
        return max([own] + [1 + depth(k, seen + (n,)) for k in kids if k not in seen])

    return max(depth(p["product_id"]) for p in E["Product"])


def check_uci_allocation(E: dict, paths: Paths, C: Checks) -> None:
    S = "public_data"
    sales = E["SalesHistory"]
    C.add(S, "SalesHistory 전부 SYNTHETIC/ALLOCATED", all(s["origin_class"] == "SYNTHETIC" and s["generation_method"] == "ALLOCATED" for s in sales))
    if not paths.uci_csv.exists():
        C.add(S, "UCI 원본 존재", False, "원본 없음 — 합계 보존 검증 불가 (미검증)")
        return
    raw = paths.uci_csv.read_bytes().decode("ascii")
    rows = list(csv.reader(io.StringIO(raw), delimiter=";"))[1:]
    stg_path = paths.root / "data" / "raw" / "public" / "uci_daily_demand" / "staging_uci_daily_orders_v1.json"
    stg = json.loads(stg_path.read_text(encoding="utf-8"))
    uid_to_row = {r["record_uid"]: r["data_row"] for r in stg["records"]}
    C.add(S, "스테이징 레코드 60 = 원본 행 수", len(uid_to_row) == len(rows) == 60, len(rows))
    C.add(S, "스테이징 값 = 원본 값 (Target 열)", all(r["target_total"] == rows[r["data_row"] - 1][12] for r in stg["records"]))
    groups = defaultdict(list)
    for s in sales:
        groups[s["derived_from_record_uid"]].append(s)
    bad_sum, bad_share, bad_code = [], [], []
    for uid, g in groups.items():
        target = dec(rows[uid_to_row[uid] - 1][12])
        if sum((dec(s["sales_qty"]) for s in g), Z) != target:
            bad_sum.append(uid)
        if sum((dec(s["allocation_share"]) for s in g), Z) != 1:
            bad_share.append(uid)
        if any(s["source_day_of_week_code"] != int(rows[uid_to_row[uid] - 1][1]) for s in g):
            bad_code.append(uid)
    C.add(S, "원본 행마다 배분 3건", len(groups) == 60 and all(len(g) == 3 for g in groups.values()), len(groups))
    C.add(S, "배분 합계 = 원본 Target (정확 일치)", not bad_sum, bad_sum)
    C.add(S, "배분 비율 합 = 1", not bad_share, bad_share)
    C.add(S, "원본 요일 코드를 숫자 그대로 보존", not bad_code, bad_code)
    C.add(S, "모든 판매 날짜가 합성값으로 표시됨 (date_is_synthetic)", all(s["date_is_synthetic"] is True for s in sales))
    dates = sorted({s["period_start"] for s in sales})
    C.add(S, "합성 날짜 = 원본 행 순서의 연속 영업일 60일 (요일 코드 해석 미사용)",
          len(dates) == 60 and all(day(x).weekday() < 5 for x in dates)
          and all(sum(1 for i in range(1, (day(b) - day(a)).days + 1) if (day(a) + dt.timedelta(days=i)).weekday() < 5) == 1
                  for a, b in zip(dates, dates[1:])), [dates[0], dates[-1]])


# ===========================================================================
# 2. Vendor 파일 대조
# ===========================================================================
def read_csv_rows(path: Path, encoding: str) -> list[list[str]]:
    data = path.read_bytes()
    if encoding == "utf-8-sig" and not data.startswith(b"\xef\xbb\xbf"):
        raise ValueError(f"{path.name}: UTF-8 BOM 없음")
    if encoding == "utf-8" and data.startswith(b"\xef\xbb\xbf"):
        raise ValueError(f"{path.name}: 예상치 못한 BOM")
    return list(csv.reader(io.StringIO(data.decode(encoding), newline="")))


def read_sheet(path: Path, sheet: str) -> list[list]:
    wb = openpyxl.load_workbook(path, read_only=False)
    ws = wb[sheet]
    return [[c.value for c in row] for row in ws.iter_rows()]


def bnum(s) -> Decimal | None:
    if s is None or s == "":
        return None
    return Decimal(str(s).replace(",", ""))


def bdate(s) -> str | None:
    if s in (None, ""):
        return None
    y, m, d = str(s).split(".")
    return dt.date(int(y), int(m), int(d)).isoformat()


def check_vendors(gt: dict, paths: Paths, C: Checks) -> None:
    E = gt["entities"]
    S = "vendor_files"
    cmap = read_json(paths.column_map)
    with open(paths.id_map, encoding="utf-8", newline="") as f:
        idm = list(csv.DictReader(f))
    to_c = {(r["vendor"], r["entity"], r["vendor_code"]): r["canonical_id"] for r in idm}
    unit = {u["unit_code"]: u for u in E["Unit"]}
    base_of = {**{p["product_id"]: p["base_unit"] for p in E["Product"]}, **{m["material_id"]: m["base_unit"] for m in E["Material"]}}

    def in_base(qty: Decimal | None, u: str, item: str) -> Decimal | None:
        if qty is None:
            return None
        return qty * dec(unit[u]["to_base_factor"]) / dec(unit[base_of[item]]["to_base_factor"])

    def cid(v, ent, code):
        return to_c[(v, ent, str(code))]

    def item_ent(code_type):
        return "Material" if code_type in ("MATERIAL",) else "Product"

    canon = {
        "Product": {p["product_id"]: p for p in E["Product"]}, "Material": {m["material_id"]: m for m in E["Material"]},
        "Supplier": {s["supplier_id"]: s for s in E["Supplier"]},
        "SupplierMaterial": {(s["supplier_id"], s["material_id"]): s for s in E["SupplierMaterial"]},
        "BOM": {b["bom_id"]: b for b in E["BOM"]}, "BOMLine": {(l["bom_id"], l["line_no"]): l for l in E["BOMLine"]},
        "Order": {o["order_id"]: o for o in E["Order"]},
        "Inventory": {(i["item_type"], i["item_id"]): i for i in E["Inventory"]},
        "Resource": {r["resource_id"]: r for r in E["Resource"]}, "Operation": {o["operation_id"]: o for o in E["Operation"]},
    }

    # canonical 기준값 (비교용, 기준 단위)
    def c_product(p):
        return (p["product_name"], p["product_type"], p["base_unit"], dec(p["safety_stock_qty"]), p["manuf_lead_time_days"],
                p["outsourcing_allowed"], dec(p["outsourcing_unit_cost"]) if p["outsourcing_unit_cost"] else None)

    def c_material(m):
        return (m["material_name"], m["base_unit"], dec(m["unit_cost"]), dec(m["safety_stock_qty"]))

    def c_bomline(l):
        comp = l["component_material_id"] or l["component_product_id"]
        return (l["component_type"], comp, in_base(dec(l["qty_per"]), l["unit"], comp), dec(l["scrap_rate"]))

    def c_bom(b):
        return (b["product_id"], b["bom_version"], b["status"], b["is_active"], b["valid_from"], b["valid_to"],
                dec(b["min_qty"]), dec(b["max_qty"]) if b["max_qty"] else None, b["site_id"], dec(b["base_qty"]))

    def c_inv(i):
        return tuple(in_base(dec(i[f]), i["unit"], i["item_id"]) for f in ("on_hand_qty", "allocated_qty", "on_order_qty")) + \
            (i["expected_receipt_date"], i["as_of_date"])

    def c_order(o):
        return (o["product_id"], in_base(dec(o["order_qty"]), o["unit"], o["product_id"]), o["order_date"], o["due_date"],
                o["priority"], o["status"])

    def c_sm(s):
        return (s["lead_time_days"], dec(s["min_order_qty"]) if s["min_order_qty"] else None, s["is_preferred"])

    found = defaultdict(lambda: defaultdict(set))
    mismatches = []

    def compare(vendor, file, ent, key, got, want):
        found[(vendor, file)][ent].add(key)
        if got != want:
            mismatches.append({"vendor": vendor, "file": file, "entity": ent, "key": str(key), "got": str(got), "want": str(want)})

    # ---- Vendor A ----
    A = paths.vendor_a
    yn = {"Y": True, "N": False}

    def a_rows(name):
        spec = next(f for f in cmap["files"] if f["file"] == f"vendor_a/{name}")
        rows = read_csv_rows(A / name, spec["encoding"])
        C.add(S, f"A {name} 헤더 = 열 대응표", rows[0] == spec["columns"], rows[0])
        return [dict(zip(rows[0], r)) for r in rows[1:]]

    def a_product(r):
        return (r["product_name"], {"FG": "FINISHED", "SF": "SEMI_FINISHED"}[r["product_type"]], r["uom"], dec(r["safety_stock"]),
                int(r["mfg_lead_time_days"]), yn[r["outsourcing_allowed"]], bnum(r["outsourcing_unit_cost_krw"]))

    def a_material(r):
        return (r["material_name"], r["uom"], dec(r["unit_cost_krw"]), dec(r["safety_stock"]))

    def a_bom(r, file):
        b = cid("A", "BOM", r["bom_code"])
        comp_t = {"MAT": "MATERIAL", "PRD": "PRODUCT"}[r["component_type"]]
        comp = cid("A", "Material" if comp_t == "MATERIAL" else "Product", r["component_code"])
        compare("A", file, "BOMLine", (b, int(r["line_no"])), (comp_t, comp, in_base(dec(r["qty_per"]), r["uom"], comp), dec(r["scrap_rate"])),
                c_bomline(canon["BOMLine"][(b, int(r["line_no"]))]))
        compare("A", file, "BOM", b, (cid("A", "Product", r["product_code"]), r["bom_version"], r["status"], yn[r["active"]],
                                      r["valid_from"], r["valid_to"] or None, dec(r["min_qty"]), bnum(r["max_qty"]), r["site"] or None,
                                      dec(r["base_qty"])), c_bom(canon["BOM"][b]))

    for r in a_rows("products.csv"):
        k = cid("A", "Product", r["product_code"])
        compare("A", "products.csv", "Product", k, a_product(r), c_product(canon["Product"][k]))
    for r in a_rows("materials.csv"):
        k = cid("A", "Material", r["material_code"])
        compare("A", "materials.csv", "Material", k, a_material(r), c_material(canon["Material"][k]))
    for r in a_rows("suppliers.csv"):
        k = cid("A", "Supplier", r["supplier_code"])
        compare("A", "suppliers.csv", "Supplier", k, r["supplier_name"], canon["Supplier"][k]["supplier_name"])
    for r in a_rows("supplier_materials.csv"):
        k = (cid("A", "Supplier", r["supplier_code"]), cid("A", "Material", r["material_code"]))
        compare("A", "supplier_materials.csv", "SupplierMaterial", k, (int(r["lead_time_days"]), bnum(r["moq"]), yn[r["preferred"]]),
                c_sm(canon["SupplierMaterial"][k]))
    for r in a_rows("bom.csv"):
        a_bom(r, "bom.csv")
    for r in a_rows("inventory.csv"):
        t = {"PRO": "PRODUCT", "MAT": "MATERIAL"}[r["item_type"]]
        k = (t, cid("A", "Material" if t == "MATERIAL" else "Product", r["item_code"]))
        got = tuple(in_base(dec(r[f]), r["uom"], k[1]) for f in ("on_hand", "allocated", "on_order")) + (r["expected_receipt"] or None, r["as_of"])
        compare("A", "inventory.csv", "Inventory", k, got, c_inv(canon["Inventory"][k]))
    for r in a_rows("orders.csv"):
        k = cid("A", "Order", r["order_no"])
        p = cid("A", "Product", r["product_code"])
        compare("A", "orders.csv", "Order", k, (p, in_base(dec(r["qty"]), r["uom"], p), r["order_date"], r["due_date"], r["priority"], r["status"]),
                c_order(canon["Order"][k]))
    for r in a_rows("resources.csv"):
        k = cid("A", "Resource", r["resource_code"])
        c = canon["Resource"][k]
        compare("A", "resources.csv", "Resource", k,
                (r["resource_name"], r["resource_type"], dec(r["regular_hours_per_day"]), dec(r["max_overtime_hours_per_day"]),
                 dec(r["cost_per_hour_krw"]), dec(r["overtime_cost_per_hour_krw"])),
                (c["resource_name"], c["resource_type"], dec(c["regular_hours_per_day"]), dec(c["overtime_hours_per_day_max"]),
                 dec(c["cost_per_hour"]), dec(c["overtime_cost_per_hour"])))
    for r in a_rows("routing.csv"):
        c = canon["Operation"][r["operation_code"]]
        compare("A", "routing.csv", "Operation", r["operation_code"],
                (cid("A", "Product", r["product_code"]), int(r["seq"]), cid("A", "Resource", r["resource_code"]),
                 dec(r["run_time_h_per_unit"]), dec(r["setup_time_h"])),
                (c["product_id"], c["sequence_no"], c["resource_id"], dec(c["run_time_per_unit_h"]), dec(c["setup_time_h"])))
    xa = A / "vendor_a_master.xlsx"
    for sheet, fn in (("Products", "p"), ("Materials", "m"), ("BOM", "b")):
        rows = read_sheet(xa, sheet)
        spec = next(f for f in cmap["files"] if f["file"] == "vendor_a/vendor_a_master.xlsx" and f["sheet"] == sheet)
        C.add(S, f"A xlsx {sheet} 헤더 = 열 대응표", rows[0] == spec["columns"], rows[0])
        recs = [dict(zip(rows[0], ["" if v is None else str(v) for v in r])) for r in rows[1:]]
        lead0 = all(isinstance(v, str) for r in rows[1:] for v in r[:1])
        C.add(S, f"A xlsx {sheet} 코드 열이 텍스트 (선행 0 보존)", lead0)
        for r in recs:
            if fn == "p":
                k = cid("A", "Product", r["product_code"])
                compare("A", f"xlsx:{sheet}", "Product", k, a_product(r), c_product(canon["Product"][k]))
            elif fn == "m":
                k = cid("A", "Material", r["material_code"])
                compare("A", f"xlsx:{sheet}", "Material", k, a_material(r), c_material(canon["Material"][k]))
            else:
                a_bom(r, f"xlsx:{sheet}")

    # ---- Vendor B ----
    B = paths.vendor_b
    xb = B / "vendor_b_master.xlsx"
    wb = openpyxl.load_workbook(xb)
    ko = {"예": True, "아니오": False}
    ulab = lambda u: "EA" if u == "개" else u
    for sheet in ("제품", "자재", "공급조건", "BOM"):
        ws = wb[sheet]
        spec = next(f for f in cmap["files"] if f["file"] == "vendor_b/vendor_b_master.xlsx" and f["sheet"] == sheet)
        hdr = [c.value for c in ws[2]]
        merged = [str(m) for m in ws.merged_cells.ranges]
        C.add(S, f"B xlsx {sheet} 2행 헤더 = 열 대응표", hdr == spec["columns"], hdr)
        C.add(S, f"B xlsx {sheet} 1행 병합 헤더 존재", len(merged) >= 1 and all(m.endswith("1") for m in merged), merged)
        recs = [dict(zip(hdr, r)) for r in ([c.value for c in row] for row in ws.iter_rows(min_row=3))]
        for r in recs:
            if sheet == "제품":
                k = cid("B", "Product", r["품목코드"])
                got = (r["품목명"], {"완제품": "FINISHED", "반제품": "SEMI_FINISHED"}[r["구분"]], ulab(r["단위"]), bnum(r["안전재고"]),
                       int(r["생산리드타임(일)"]), ko[r["외주가능"]], bnum(r["외주단가(원)"]))
                compare("B", f"xlsx:{sheet}", "Product", k, got, c_product(canon["Product"][k]))
            elif sheet == "자재":
                k = cid("B", "Material", r["자재코드"])
                u = ulab(r["단위"])
                base = base_of[k]
                cost = bnum(r["단가(원/단위)"]) * dec(unit[base]["to_base_factor"]) / dec(unit[u]["to_base_factor"])
                got = (r["자재명"], base, cost, in_base(bnum(r["안전재고"]), u, k))
                C.add(S, f"B 자재 {k} 표시 단위가 g(질량)/EA", u == ("g" if unit[base]["dimension"] == "MASS" else "EA"), u)
                compare("B", f"xlsx:{sheet}", "Material", k, got, c_material(canon["Material"][k]))
            elif sheet == "공급조건":
                s, mm = cid("B", "Supplier", r["공급사코드"]), cid("B", "Material", r["자재코드"])
                moq = bnum(r["최소주문량"])
                if moq is not None and unit[base_of[mm]]["dimension"] == "MASS":
                    moq = in_base(moq, "g", mm)
                compare("B", f"xlsx:{sheet}", "SupplierMaterial", (s, mm), (int(r["리드타임(일)"]), moq, ko[r["우선공급"]]),
                        c_sm(canon["SupplierMaterial"][(s, mm)]))
                compare("B", f"xlsx:{sheet}", "Supplier", s, r["공급사명"], canon["Supplier"][s]["supplier_name"])
            else:
                b = cid("B", "BOM", r["BOM번호"])
                ct = {"자재": "MATERIAL", "제품": "PRODUCT"}[r["자품목구분"]]
                comp = cid("B", "Material" if ct == "MATERIAL" else "Product", r["자품목코드"])
                compare("B", f"xlsx:{sheet}", "BOMLine", (b, int(r["순번"])),
                        (ct, comp, in_base(bnum(r["소요량"]), ulab(r["단위"]), comp), bnum(r["손실률(%)"]) / 100),
                        c_bomline(canon["BOMLine"][(b, int(r["순번"]))]))
                compare("B", f"xlsx:{sheet}", "BOM", b,
                        (cid("B", "Product", r["모품목코드"]), r["버전"], {"승인": "APPROVED", "초안": "DRAFT"}[r["상태"]], ko[r["활성"]],
                         bdate(r["시작일"]), bdate(r["종료일"]), bnum(r["최소수량"]), bnum(r["최대수량"]), r["사이트"] or None,
                         bnum(r["기준수량"])), c_bom(canon["BOM"][b]))
    st = {"진행": "OPEN", "완료": "CLOSED", "취소": "CANCELLED"}
    for name, ent in (("vendor_b_orders.csv", "Order"), ("vendor_b_inventory.csv", "Inventory")):
        raw = (B / name).read_bytes()
        C.add(S, f"B {name} CP949 (UTF-8로는 해독 불가)", _not_utf8(raw))
        C.add(S, f"B {name} 줄바꿈 CRLF", raw.count(b"\r\n") == raw.count(b"\n"))
        rows = read_csv_rows(B / name, "cp949")
        recs = [dict(zip(rows[0], r)) for r in rows[1:]]
        for r in recs:
            if ent == "Order":
                k = cid("B", "Order", r["주문번호"])
                p = cid("B", "Product", r["품목코드"])
                compare("B", name, ent, k, (p, in_base(bnum(r["수량"]), ulab(r["단위"]), p), bdate(r["주문일자"]), bdate(r["납기일자"]),
                                            {"긴급": "URGENT", "일반": "NORMAL"}[r["긴급여부"]], st[r["상태"]]), c_order(canon["Order"][k]))
            else:
                t = {"제품": "PRODUCT", "자재": "MATERIAL"}[r["구분"]]
                k = (t, cid("B", "Material" if t == "MATERIAL" else "Product", r["품목코드"]))
                u = ulab(r["단위"])
                got = tuple(in_base(bnum(r[f]), u, k[1]) for f in ("현재고", "할당", "입고예정")) + (bdate(r["입고예정일"]), bdate(r["기준일"]))
                compare("B", name, ent, k, got, c_inv(canon["Inventory"][k]))
    C.add(S, "Vendor 값 = canonical 업무 의미 (불일치 0)", not mismatches, mismatches[:20])
    expected_cover = {
        ("A", "products.csv"): {"Product": set(canon["Product"])}, ("A", "materials.csv"): {"Material": set(canon["Material"])},
        ("A", "bom.csv"): {"BOMLine": set(canon["BOMLine"]), "BOM": set(canon["BOM"])},
        ("A", "inventory.csv"): {"Inventory": set(canon["Inventory"])}, ("A", "orders.csv"): {"Order": set(canon["Order"])},
        ("A", "routing.csv"): {"Operation": set(canon["Operation"])},
        ("A", "supplier_materials.csv"): {"SupplierMaterial": set(canon["SupplierMaterial"])},
        ("B", "xlsx:BOM"): {"BOMLine": set(canon["BOMLine"]), "BOM": set(canon["BOM"])},
        ("B", "vendor_b_orders.csv"): {"Order": set(canon["Order"])}, ("B", "vendor_b_inventory.csv"): {"Inventory": set(canon["Inventory"])},
        ("B", "xlsx:공급조건"): {"SupplierMaterial": set(canon["SupplierMaterial"])},
    }
    gaps = [(f, e) for f, ents in expected_cover.items() for e, keys in ents.items() if found[f][e] != keys]
    C.add(S, "Vendor 파일이 canonical 레코드를 빠짐없이 표현", not gaps, gaps)
    encs = {f["encoding"] for f in cmap["files"] if f["format"] == "csv"}
    C.add(S, "CSV 인코딩 3종 포함 (utf-8, utf-8-sig, cp949)", encs == {"utf-8", "utf-8-sig", "cp949"}, sorted(encs))
    # 대응표: 행 번호가 실제로 해당 레코드를 가리키는가
    with open(paths.row_map, encoding="utf-8", newline="") as f:
        rmap = list(csv.DictReader(f))
    uid_of = {r["record_uid"]: (ent, r) for ent, rows in E.items() for r in rows}
    C.add(S, "대응표의 canonical_record_uid가 모두 존재", all(r["canonical_record_uid"] in uid_of for r in rmap), len(rmap))
    bad = []
    for r in rmap:
        if r["entity"] not in ("Order", "Product", "Material", "Operation"):
            continue
        rows = read_sheet(paths.root / "data" / r["file"], r["sheet"]) if r["sheet"] else read_csv_rows(
            paths.root / "data" / r["file"], next(f["encoding"] for f in cmap["files"] if f["file"] == r["file"]))
        code = rows[int(r["row_number"]) - 1][0]
        ent_key = {"Order": "Order", "Product": "Product", "Material": "Material"}.get(r["entity"])
        mapped = cid(r["vendor"], ent_key, code) if ent_key else str(code)
        if mapped != r["canonical_key"]:
            bad.append(r)
    C.add(S, "대응표 행 번호 → 실제 행의 코드가 canonical 키와 일치 (Order/Product/Material/Operation)", not bad, bad[:5])


def _not_utf8(raw: bytes) -> bool:
    try:
        raw.decode("utf-8")
        return False
    except UnicodeDecodeError:
        return True


# ===========================================================================
# 3. 오류 주입 확인
# ===========================================================================
def load_any(root: Path, rel: str, sheet: str, cmap: dict) -> tuple[list[list], int]:
    path = root / "data" / rel
    if rel.endswith(".xlsx"):
        rows = read_sheet(path, sheet)
        spec = next(f for f in cmap["files"] if f["file"] == rel.replace("invalid/", "") and f.get("sheet") == sheet)
        return rows, spec["header_rows"]
    spec = next(f for f in cmap["files"] if f["file"] == rel.replace("invalid/", ""))
    return read_csv_rows(path, spec["encoding"]), 1


def check_injection(gt: dict, paths: Paths, C: Checks) -> dict:
    S = "error_injection"
    inj = read_json(paths.injected)
    cmap = read_json(paths.column_map)
    errs = inj["errors"]
    C.add(S, "주입 목록 13건·12유형", len(errs) == 13 and len({e["error_code"] for e in errs}) == 12,
          {"count": len(errs), "codes": sorted({e["error_code"] for e in errs})})
    C.add(S, "오류 ID 유일", len({e["error_id"] for e in errs}) == len(errs))
    files = sorted({(e["clean_file"], e["sheet"]) for e in errs})
    # 오류가 선언되지 않은 파일·시트는 정상본과 값이 같아야 한다
    all_targets = [(f["file"], f.get("sheet", "")) for f in cmap["files"]]
    untouched_diff = []
    for rel, sheet in all_targets:
        if (rel, sheet) in files:
            continue
        a, _ = load_any(paths.root, rel, sheet, cmap)
        b, _ = load_any(paths.root, "invalid/" + rel, sheet, cmap)
        if a != b:
            untouched_diff.append((rel, sheet))
    C.add(S, "오류 미선언 파일·시트는 정상본과 동일", not untouched_diff, untouched_diff)
    confirmed = []
    for rel, sheet in files:
        clean, hrows = load_any(paths.root, rel, sheet, cmap)
        bad, _ = load_any(paths.root, "invalid/" + rel, sheet, cmap)
        header = clean[hrows - 1]
        diffs = set()
        for r in range(max(len(clean), len(bad))):
            if r >= len(clean):
                diffs.add(("append", r + 1))
                continue
            if r >= len(bad):
                diffs.add(("removed", r + 1))
                continue
            for c in range(max(len(clean[r]), len(bad[r]))):
                x = clean[r][c] if c < len(clean[r]) else None
                y = bad[r][c] if c < len(bad[r]) else None
                if (x or None) != (y or None):
                    diffs.add(("cell", r + 1, c + 1))
        declared = set()
        for e in (e for e in errs if (e["clean_file"], e["sheet"]) == (rel, sheet)):
            for ch in e["changes"]:
                if ch.get("append"):
                    declared.add(("append", ch.get("row", e["row_number"])))
                elif "column_index" in ch:
                    declared.add(("cell", ch["row"], ch["column_index"]))
                else:
                    declared.add(("cell", ch["row"], header.index(ch["column"]) + 1))
        C.add(S, f"{rel} {sheet}: 실제 차이 = 선언된 주입 위치", diffs == declared,
              {"undeclared": sorted(map(str, diffs - declared)), "not_found": sorted(map(str, declared - diffs))})
    for e in errs:
        ok, why = predicate(e, paths, cmap)
        C.add(S, f"{e['error_id']} {e['error_code']} 판정 조건 (주입본 참·정상본 거짓)", ok, why)
        if ok:
            confirmed.append(e["error_id"])
    C.add(S, "주입 개수 = 독립 확인 개수", len(confirmed) == len(errs), {"injected": len(errs), "confirmed": len(confirmed)})
    return {"injected": len(errs), "confirmed": len(confirmed), "confirmed_ids": confirmed}


def predicate(e: dict, paths: Paths, cmap: dict) -> tuple[bool, str]:
    def rows_of(which):
        rel = e["clean_file"] if which == "clean" else e["file"]
        r, h = load_any(paths.root, rel, e["sheet"], cmap)
        return r, h

    def check(which):
        rows, h = rows_of(which)
        hdr = rows[h - 1]
        row = rows[e["row_number"] - 1] if e["row_number"] - 1 < len(rows) else None
        get = lambda col: (row[hdr.index(col)] if row is not None else None)
        code = e["error_code"]
        if code == "ERR-MISSING":
            return get(e["field"]) in ("", None)
        if code == "ERR-DUPID":
            col = [r[hdr.index(e["field"])] for r in rows[h:]]
            return Counter(col)[get(e["field"])] >= 2 if row else False
        if code == "ERR-BADREF":
            mats, _ = load_any(paths.root, e["file"] if which == "invalid" else e["clean_file"], "자재", cmap)
            return get("자품목구분") == "자재" and get(e["field"]) not in {r[0] for r in mats[2:]}
        if code == "ERR-UNIT":
            mats = read_csv_rows(paths.root / "data" / (e["file"] if which == "invalid" else e["clean_file"]).replace("bom.csv", "materials.csv"), "utf-8-sig")
            mu = {r[0]: r[2] for r in mats[1:]}
            dims = {"kg": "MASS", "g": "MASS", "EA": "COUNT", "h": "TIME"}
            return get("component_type") == "MAT" and dims[get("uom")] != dims[mu[get("component_code")]]
        if code == "ERR-NUMFMT":
            try:
                Decimal(str(get(e["field"])).replace(",", ""))
                return False
            except InvalidOperation:
                return True
        if code == "ERR-TABLE":
            spec = next(f for f in cmap["files"] if f["file"] == e["clean_file"] and f.get("sheet") == e["sheet"])
            return hdr != spec["columns"]
        if code == "ERR-IDVER":
            same = [r for r in rows[h:] if r[hdr.index("bom_code")] == get("bom_code")]
            return len({(r[hdr.index("product_code")], r[hdr.index("bom_version")]) for r in same}) > 1
        if code == "ERR-DATE":
            if e["vendor"] == "A":
                try:
                    dt.date.fromisoformat(get("due_date"))
                    return dt.date.fromisoformat(get("due_date")) < dt.date.fromisoformat(get("order_date"))
                except ValueError:
                    return True
            return bdate(get("납기일자")) < bdate(get("주문일자"))
        if code == "ERR-BOMCYCLE":
            g = defaultdict(set)
            for r in rows[h:]:
                if r[hdr.index("자품목구분")] == "제품":
                    g[r[hdr.index("모품목코드")]].add(r[hdr.index("자품목코드")])
            return bool(find_cycles(g))
        if code == "ERR-BOMVER":
            if row is None:
                return False
            prod = get("product_code")
            vers = {}
            for r in rows[h:]:
                d = dict(zip(hdr, r))
                if d["product_code"] == prod and d["status"] == "APPROVED" and d["active"] == "Y":
                    vers[d["bom_code"]] = d
            far, inf = "9999-12-31", Decimal("1e30")
            vs = list(vers.values())
            for i, x in enumerate(vs):
                for y in vs[i + 1:]:
                    if (x["valid_from"] <= (y["valid_to"] or far) and y["valid_from"] <= (x["valid_to"] or far)
                            and dec(x["min_qty"]) < (bnum(y["max_qty"]) or inf) and dec(y["min_qty"]) < (bnum(x["max_qty"]) or inf)
                            and (not x["site"] or not y["site"] or x["site"] == y["site"])):
                        return True
            return False
        if code == "ERR-INVQTY":
            return dec(get("allocated")) > dec(get("on_hand"))
        if code == "ERR-RESREL":
            res = read_csv_rows(paths.root / "data" / (e["file"] if which == "invalid" else e["clean_file"]).replace("routing.csv", "resources.csv"), "utf-8")
            return get("resource_code") not in {r[0] for r in res[1:]}
        raise ValueError(code)

    inv_ok, clean_ok = check("invalid"), check("clean")
    return inv_ok and not clean_ok, f"invalid={inv_ok}, clean={clean_ok}"


# ===========================================================================
# 4. 기대값 재계산 (구현 B)
# ===========================================================================
class Calc:
    def __init__(self, gt: dict) -> None:
        self.E, self.P = gt["entities"], gt["parameters"]
        self.factor = {u["unit_code"]: (u["dimension"], dec(u["to_base_factor"])) for u in self.E["Unit"]}
        self.item = {}
        for p in self.E["Product"]:
            self.item[("PRODUCT", p["product_id"])] = p
        for m in self.E["Material"]:
            self.item[("MATERIAL", m["material_id"])] = m
        self.ref = day(self.P["reference_date"])
        self.hs, self.he = day(self.P["horizon_start"]), day(self.P["horizon_end"])
        self.workdays = [self.ref + dt.timedelta(days=i) for i in range(1, 400) if (self.ref + dt.timedelta(days=i)).weekday() < 5]

    def conv(self, qty: Decimal, unit: str, key: tuple) -> Decimal:
        d1, f1 = self.factor[unit]
        d2, f2 = self.factor[self.item[key]["base_unit"]]
        assert d1 == d2, (unit, key)
        return qty * f1 / f2

    def inv(self, key: tuple, horizon: bool = True) -> Decimal:
        recs = [i for i in self.E["Inventory"] if (i["item_type"], i["item_id"]) == key]
        total = Z
        for i in recs:
            total += self.conv(dec(i["on_hand_qty"]) - dec(i["allocated_qty"]), i["unit"], key)
            if horizon and dec(i["on_order_qty"]) > 0 and day(i["expected_receipt_date"]) <= self.he:
                total += self.conv(dec(i["on_order_qty"]), i["unit"], key)
        return total

    def pick_bom(self, prod: str, d: dt.date, q: Decimal) -> list[str]:
        def ok(b):
            return (b["product_id"] == prod and b["is_active"] and b["status"] == "APPROVED"
                    and day(b["valid_from"]) <= d <= (day(b["valid_to"]) if b["valid_to"] else dt.date.max)
                    and dec(b["min_qty"]) <= q and (not b["max_qty"] or q < dec(b["max_qty"]))
                    and (b["site_id"] in (None, self.P["site_id"])))
        return sorted(b["bom_id"] for b in self.E["BOM"] if ok(b))

    def comps(self, bom_id: str) -> list[tuple]:
        base = next(dec(b["base_qty"]) for b in self.E["BOM"] if b["bom_id"] == bom_id)
        out = []
        for l in sorted((l for l in self.E["BOMLine"] if l["bom_id"] == bom_id), key=lambda l: l["line_no"]):
            key = ("MATERIAL", l["component_material_id"]) if l["component_material_id"] else ("PRODUCT", l["component_product_id"])
            out.append((l["line_no"], key, self.conv(dec(l["qty_per"]), l["unit"], key) * (1 + dec(l["scrap_rate"])) / base))
        return out

    def open_orders(self):
        return [o for o in self.E["Order"] if o["status"] == "OPEN" and self.hs <= day(o["due_date"]) <= self.he]

    def ops(self, prod: str):
        return sorted((o for o in self.E["Operation"] if o["product_id"] == prod), key=lambda o: o["sequence_no"])

    def res(self, rid: str):
        return next(r for r in self.E["Resource"] if r["resource_id"] == rid)

    # --- 위상 정렬 기반 다단계 순소요 ---
    def requirements(self):
        prods = [p["product_id"] for p in self.E["Product"]]
        indep = {p: sum((self.conv(dec(o["order_qty"]), o["unit"], ("PRODUCT", p)) for o in self.open_orders() if o["product_id"] == p), Z)
                 for p in prods}
        parents = defaultdict(set)
        for b in self.E["BOM"]:
            if b["is_active"] and b["status"] == "APPROVED":
                for l in self.E["BOMLine"]:
                    if l["bom_id"] == b["bom_id"] and l["component_product_id"]:
                        parents[l["component_product_id"]].add(b["product_id"])
        done, order = set(), []
        queue = deque(sorted(p for p in prods if not parents[p]))
        while queue:
            p = queue.popleft()
            order.append(p)
            done.add(p)
            for q in sorted(prods):
                if q not in done and q not in queue and parents[q] and parents[q] <= done:
                    queue.append(q)
        assert len(order) == len(prods), "순환 또는 미처리 품목"
        dep = defaultdict(lambda: Z)
        net, chosen, gross_mat, explosion = {}, {}, defaultdict(lambda: Z), []
        for p in order:
            ss = dec(self.item[("PRODUCT", p)]["safety_stock_qty"]) if self.P["include_safety_stock"] else Z
            n = indep[p] + dep[p] + ss - self.inv(("PRODUCT", p))
            net[p] = n if n > 0 else Z
            if net[p] == 0:
                chosen[p] = None
                continue
            c = self.pick_bom(p, self.hs, net[p])
            assert len(c) == 1, (p, c)
            chosen[p] = c[0]
            for ln, key, per in self.comps(c[0]):
                explosion.append((p, c[0], ln, key[1], net[p] * per))
                if key[0] == "PRODUCT":
                    dep[key[1]] += net[p] * per
                else:
                    gross_mat[key[1]] += net[p] * per
        short = {}
        for m in self.E["Material"]:
            mid = m["material_id"]
            ss = dec(m["safety_stock_qty"]) if self.P["include_material_safety_stock"] else Z
            s = gross_mat[mid] + ss - self.inv(("MATERIAL", mid))
            short[mid] = s if s > 0 else Z
        return {"net": net, "bom": chosen, "dep": dict(dep), "gross": dict(gross_mat), "short": short, "explosion": explosion}

    def wd_add(self, start: dt.date, n: int) -> dt.date:
        if n == 0:
            return start
        after = [d for d in self.workdays if d > start] if start >= self.ref else None
        return after[n - 1]

    def wd_count(self, until: dt.date) -> int:
        return sum(1 for d in self.workdays if d <= until)

    def feasibility(self, R, lt_change=None):
        lt = {s["material_id"]: s["lead_time_days"] for s in self.E["SupplierMaterial"] if s["is_preferred"]}
        lt.update(lt_change or {})
        short = {k for k, v in R["short"].items() if v > 0}
        left = {p["product_id"]: self.inv(("PRODUCT", p["product_id"]), horizon=False) for p in self.E["Product"]}
        used = defaultdict(lambda: Z)
        out = {}
        for o in sorted(self.open_orders(), key=lambda o: (o["due_date"], o["order_id"])):
            p, due = o["product_id"], day(o["due_date"])
            need = self.conv(dec(o["order_qty"]), o["unit"], ("PRODUCT", p))
            use = need if left[p] >= need else left[p]
            left[p] -= use
            make = need - use
            hrs = defaultdict(lambda: Z)
            mats, lt_days = set(), 0
            if make > 0:
                bom = R["bom"][p]
                for op in self.ops(p):
                    hrs[op["resource_id"]] += make * dec(op["run_time_per_unit_h"]) + dec(op["setup_time_h"])
                extra = []
                for _, key, per in self.comps(bom):
                    if key[0] == "MATERIAL":
                        mats.add(key[1])
                        continue
                    req = make * per
                    have = left[key[1]] if left[key[1]] < req else req
                    left[key[1]] -= have
                    sf_make = req - have
                    if sf_make > 0:
                        extra.append(self.item[key]["manuf_lead_time_days"])
                        for op in self.ops(key[1]):
                            hrs[op["resource_id"]] += sf_make * dec(op["run_time_per_unit_h"]) + dec(op["setup_time_h"])
                        mats |= {k[1] for _, k, _ in self.comps(R["bom"][key[1]]) if k[0] == "MATERIAL"}
                lt_days = self.item[("PRODUCT", p)]["manuf_lead_time_days"] + (max(extra) if extra else 0)
            over = []
            for rid in sorted(hrs):
                used[rid] += hrs[rid]
                if used[rid] > dec(self.res(rid)["regular_hours_per_day"]) * self.wd_count(due):
                    over.append(rid)
            if make > 0:
                waits = [lt[m] for m in mats if m in short]
                ready = self.ref + dt.timedelta(days=max(waits) if waits else 0)
                need_days = 0
                for rid, h in hrs.items():
                    per_day = dec(self.res(rid)["regular_hours_per_day"])
                    dd = int((h / per_day).to_integral_value(rounding="ROUND_CEILING"))
                    need_days = max(need_days, dd)
                n = max(lt_days, need_days)
                fin = self.wd_add(ready, n)
                fin_no_wait = self.wd_add(self.ref, n)
            else:
                ready, n, fin, fin_no_wait = self.ref, 0, self.ref, self.ref
            reasons = []
            if fin > due and fin_no_wait <= due:
                reasons.append("MATERIAL_LEADTIME")
            if over:
                reasons.append("CAPACITY_EXCEEDED")
            if fin > due:
                reasons.append("DUE_DATE_MISSED")
            out[o["order_id"]] = {"from_stock": use, "production_qty": make, "material_ready_date": ready.isoformat(),
                                  "production_days": n, "expected_completion_date": fin.isoformat(),
                                  "slack_days": (due - fin).days, "capacity_exceeded_resources": over,
                                  "reasons": reasons, "feasible": not reasons}
        return out

    def plans(self):
        """계획안 평가 (구현 B): 자원·주 버킷을 먼저 만들고, 라인별 완료 수량은 제약 자원만 골라 계산."""
        res = {r["resource_id"]: r for r in self.E["Resource"]}
        weeks = [(self.hs + dt.timedelta(days=7 * k)).isoformat() for k in range(4)]
        out = {}
        for plan in self.E["ProductionPlan"]:
            ot_on = plan["overtime_allowed"]
            L = [l for l in self.E["ProductionPlanLine"] if l["plan_id"] == plan["plan_id"] and dec(l["planned_qty"]) > 0]
            need = {(r, w): Z for r in res for w in weeks}
            use, outs, out_qty = defaultdict(lambda: Z), Z, Z
            for l in L:
                q, p = dec(l["planned_qty"]), l["product_id"]
                if l["production_mode"] == "OUTSOURCE":
                    outs += q * dec(self.item[("PRODUCT", p)]["outsourcing_unit_cost"])
                    out_qty += q
                    continue
                for op in self.ops(p):
                    need[(op["resource_id"], l["period_start"])] += q * dec(op["run_time_per_unit_h"]) + dec(op["setup_time_h"])
                c = self.pick_bom(p, day(l["period_start"]), q)
                assert len(c) == 1
                for _, key, per in self.comps(c[0]):
                    if key[0] == "MATERIAL":
                        use[key[1]] += q * per
            cap = {(r, w): (dec(res[r]["regular_hours_per_day"]) * 5, dec(res[r]["overtime_hours_per_day_max"]) * 5 if ot_on else Z)
                   for r in res for w in weeks}
            reg_h = ot_h = cost_reg = cost_ot = Z
            excess = {}
            for (r, w), h in need.items():
                rc, oc = cap[(r, w)]
                used_reg = h if h < rc else rc
                spill = h - rc if h > rc else Z
                used_ot = spill if spill < oc else oc
                reg_h += used_reg
                ot_h += used_ot
                cost_reg += used_reg * dec(res[r]["cost_per_hour"])
                cost_ot += used_ot * dec(res[r]["overtime_cost_per_hour"])
                excess[(r, w)] = h - rc - oc if h > rc + oc else Z
            supply = defaultdict(list)
            unproduced = Z
            for p in self.E["Product"]:
                if p["product_type"] == "FINISHED":
                    supply[p["product_id"]].append((self.ref, self.inv(("PRODUCT", p["product_id"]), horizon=False)))
            for l in L:
                q, p, w = dec(l["planned_qty"]), l["product_id"], l["period_start"]
                fin = self.item[("PRODUCT", p)]["product_type"] == "FINISHED"
                if l["production_mode"] == "OUTSOURCE":
                    lt = [o["lead_time_days"] for o in self.P["outsourcers"].values() if o["product_id"] == p][0]
                    if fin:
                        supply[p].append((day(w) + dt.timedelta(days=lt), q))
                    continue
                tight = [(Fraction(sum(cap[(op["resource_id"], w)])), Fraction(need[(op["resource_id"], w)]))
                         for op in self.ops(p) if need[(op["resource_id"], w)] > sum(cap[(op["resource_id"], w)])]
                done = q if not tight else Decimal(math.floor(min(Fraction(q) * a / b for a, b in tight)))
                unproduced += q - done
                if fin:
                    supply[p].append((day(w) + dt.timedelta(days=4), done))
            late = []
            for p in {o["product_id"] for o in self.open_orders()}:
                running = Z
                for o in sorted((o for o in self.open_orders() if o["product_id"] == p), key=lambda o: (o["due_date"], o["order_id"])):
                    running += dec(o["order_qty"])
                    got = sum((q for d_, q in supply[p] if d_ <= day(o["due_date"])), Z)
                    if got < running:
                        late.append(o["order_id"])
            short = {}
            for m in self.E["Material"]:
                v = use[m["material_id"]] + (dec(m["safety_stock_qty"]) if self.P["include_material_safety_stock"] else Z) - self.inv(("MATERIAL", m["material_id"]))
                if v > 0:
                    short[m["material_id"]] = v
            out[plan["plan_id"]] = {"material": sum((q * dec(self.item[("MATERIAL", m)]["unit_cost"]) for m, q in use.items()), Z),
                                    "regular_production": cost_reg, "overtime": cost_ot, "outsourcing": outs,
                                    "worked_regular_h": reg_h, "worked_overtime_h": ot_h, "outsourced_qty": out_qty,
                                    "excess_total": sum(excess.values(), Z), "excess": excess,
                                    "excess_list": sorted((r, w, v) for (r, w), v in excess.items() if v > 0),
                                    "unproduced": unproduced, "shortage": short, "late": sorted(late)}
        return out


def check_expected(gt: dict, paths: Paths, C: Checks) -> dict:
    S = "decision_calc"
    X = gt["expected"]
    k = Calc(gt)
    R = k.requirements()
    eq = lambda a, b: fmt(dec(a)) == fmt(dec(b))
    prq = X["production_requirements"]
    C.add(S, "생산 필요량 (5품목) A=B", all(eq(prq[p]["net_requirement"], R["net"][p]) for p in R["net"]),
          {p: (prq[p]["net_requirement"], fmt(R["net"][p])) for p in R["net"]})
    C.add(S, "반제품 종속수요 집계 A=B", all(eq(prq[p]["dependent_demand"], R["dep"].get(p, Z)) for p in prq),
          {p: (prq[p]["dependent_demand"], fmt(R["dep"].get(p, Z))) for p in prq})
    C.add(S, "적용 BOM 버전 A=B", all(prq[p]["bom_id"] == R["bom"][p] for p in prq), {p: (prq[p]["bom_id"], R["bom"][p]) for p in prq})
    ea = sorted((e["parent"], e["bom_id"], e["line_no"], e["component_id"], fmt(dec(e["component_qty_base"]))) for e in X["bom_explosion"])
    eb = sorted((p, b, l, c, fmt(q)) for p, b, l, c, q in R["explosion"])
    C.add(S, "다단계 BOM 전개 결과 A=B", ea == eb, {"A_only": [x for x in ea if x not in eb], "B_only": [x for x in eb if x not in ea]})
    ga = {m: fmt(dec(v)) for m, v in X["material_gross_requirements"].items()}
    C.add(S, "자재별 총소요량 A=B", ga == {m: fmt(v) for m, v in R["gross"].items()}, ga)
    sa = {m: fmt(dec(v["shortage"])) for m, v in X["material_shortage"].items()}
    C.add(S, "재고 부족량 A=B", sa == {m: fmt(v) for m, v in R["short"].items()}, sa)
    # 계획 기간 자원 부하 (B: 단순 합산)
    days = sum(1 for d in k.workdays if k.hs <= d <= k.he)
    hl = X["horizon_resource_load"]
    bad = []
    for r in k.E["Resource"]:
        req = sum((R["net"][op["product_id"]] * dec(op["run_time_per_unit_h"]) + dec(op["setup_time_h"])
                   for op in k.E["Operation"] if op["resource_id"] == r["resource_id"] and R["net"][op["product_id"]] > 0), Z)
        cap = dec(r["regular_hours_per_day"]) * days
        if not (eq(hl[r["resource_id"]]["required_h"], req) and eq(hl[r["resource_id"]]["excess_h"], max(Z, req - cap))):
            bad.append(r["resource_id"])
    C.add(S, "자원별 필요시간·능력 초과 (계획 기간) A=B", not bad, bad)
    F = k.feasibility(R)
    fa = {r["order_id"]: r for r in X["order_feasibility"]}
    diffs = []
    for oid, b in F.items():
        a = fa.get(oid)
        fields = ("material_ready_date", "production_days", "expected_completion_date", "slack_days",
                  "capacity_exceeded_resources", "reasons", "feasible")
        if a is None or any(a[f] != b[f] for f in fields) or not eq(a["production_qty"], b["production_qty"]):
            diffs.append(oid)
    C.add(S, "주문별 납기 판정 (OPEN 18건) A=B", not diffs and len(F) == len(fa) == 18, {"diff": diffs, "n": len(F)})
    PB = k.plans()
    pa = X["plan_evaluation"]
    pd = []
    for pid, b in PB.items():
        a = pa[pid]
        c = a["costs_krw"]
        if not (eq(c["material"], b["material"]) and eq(c["regular_production"], b["regular_production"])
                and eq(c["overtime"], b["overtime"]) and eq(c["outsourcing"], b["outsourcing"])
                and eq(a["worked_regular_h"], b["worked_regular_h"]) and eq(a["worked_overtime_h"], b["worked_overtime_h"])
                and eq(a["capacity_excess_total_h"], b["excess_total"]) and eq(a["unproduced_qty_total"], b["unproduced"])
                and eq(sum((dec(v) for v in a["outsourced_qty"].values()), Z), b["outsourced_qty"])
                and sorted(x["order_id"] for x in a["late_orders"]) == b["late"]
                and {m: fmt(dec(v)) for m, v in a["material_shortage"].items()} == {m: fmt(v) for m, v in b["shortage"].items()}
                and sorted((x["resource_id"], x["week_start"], fmt(dec(x["excess_h"]))) for x in a["capacity"] if dec(x["excess_h"]) > 0)
                == [(r, w, fmt(v)) for r, w, v in b["excess_list"]]
                and c["inventory_holding"] == "미산정" and c["late_penalty"] == "미산정"):
            pd.append(pid)
    C.add(S, "계획안 3종 비용·근무시간·능력초과·미생산·외주량·자재부족·납기지연 A=B", not pd, pd)
    plan_sanity(k, PB, C)
    return check_cq(gt, k, R, F, PB, paths, C)


def plan_sanity(k: "Calc", PB: dict, C: Checks) -> None:
    """계획안 의미 점검: 같은 조건 적용, 잔업·외주가 능력 초과를 늘리지 않는지 (늘면 원인 기록)."""
    S = "plan_semantics"
    lines = k.E["ProductionPlanLine"]
    fg = {p["product_id"] for p in k.E["Product"] if p["product_type"] == "FINISHED"}
    by_plan = defaultdict(lambda: defaultdict(lambda: Z))
    for l in lines:
        if l["product_id"] in fg:
            by_plan[l["plan_id"]][(l["product_id"], l["period_start"])] += dec(l["planned_qty"])
    base = by_plan["PLAN-REG"]
    C.add(S, "세 안의 완제품 주별 총수량(자체+외주) 동일 (같은 주문·납기 조건)",
          all(by_plan[p] == base for p in by_plan), {p: {f"{a}|{b}": fmt(v) for (a, b), v in d.items()} for p, d in by_plan.items()})
    plans = {p["plan_id"]: p for p in k.E["ProductionPlan"]}
    C.add(S, "잔업 허용은 잔업안만", [pid for pid, p in plans.items() if p["overtime_allowed"]] == ["PLAN-OT"])
    outs = [l for l in lines if l["production_mode"] == "OUTSOURCE"]
    cap = {o["product_id"]: dec(o["capacity_per_week"]) for o in k.P["outsourcers"].values()}
    C.add(S, "외주 라인은 외주안의 외주 가능 제품만, 주별 외주처 능력 이내",
          all(l["plan_id"] == "PLAN-OUT" and k.item[("PRODUCT", l["product_id"])]["outsourcing_allowed"]
              and dec(l["planned_qty"]) <= cap[l["product_id"]] for l in outs) and bool(outs), len(outs))
    C.add(S, "외주처 조건이 합성 가정으로 표시됨", all("합성 가정" in o["name"] and o["origin_class"] == "SYNTHETIC" for o in k.P["outsourcers"].values()))
    reg = PB["PLAN-REG"]["excess"]
    for pid in ("PLAN-OT", "PLAN-OUT"):
        worse = [(r, w, fmt(reg[(r, w)]), fmt(v)) for (r, w), v in PB[pid]["excess"].items() if v > reg[(r, w)]]
        cause = None
        if worse:
            cause = ("잔업안: 같은 필요시간에 가용시간만 늘었으므로 초과가 늘 수 없음 → 계산 오류" if pid == "PLAN-OT"
                     else "외주안: 외주로 사내 필요시간이 줄어야 하므로 초과 증가는 반제품 수요 계산 오류 가능성")
        C.add(S, f"{pid} 자원·주별 능력 초과 ≤ 정규생산안", not worse, {"increased": worse, "cause": cause})
    C.add(S, "잔업안 실제 잔업시간 > 0 이고 잔업 가용 이내",
          Z < PB["PLAN-OT"]["worked_overtime_h"] <= sum(dec(r["overtime_hours_per_day_max"]) * 5 * 4 for r in k.E["Resource"]))
    C.add(S, "정규·외주안 잔업시간 = 0", PB["PLAN-REG"]["worked_overtime_h"] == 0 and PB["PLAN-OUT"]["worked_overtime_h"] == 0)


def check_cq(gt, k: Calc, R, F, PB, paths: Paths, C: Checks) -> dict:
    S = "cq"
    X = gt["expected"]["cq"]
    P = gt["parameters"]["cq"]
    status = {}

    def mark(cq, ok, detail=None):
        status[cq] = "검증 통과" if ok else "불일치"
        C.add(S, f"{cq} 기대 결과 독립 재현", ok, detail)

    # 역방향 사용처 (구현 B: 부모 맵 + 큐)
    used_by = defaultdict(set)
    for b in k.E["BOM"]:
        if b["is_active"] and b["status"] == "APPROVED":
            for l in k.E["BOMLine"]:
                if l["bom_id"] == b["bom_id"]:
                    used_by[l["component_material_id"] or l["component_product_id"]].add(b["product_id"])

    def finished_users(item):
        seen, q = set(), deque([item])
        while q:
            for par in used_by[q.popleft()]:
                if par not in seen:
                    seen.add(par)
                    q.append(par)
        return sorted(p for p in seen if k.item[("PRODUCT", p)]["product_type"] == "FINISHED")

    mark("CQ-01", {pid: sorted((x["resource_id"], x["week_start"], fmt(dec(x["excess_h"]))) for x in v) for pid, v in X["CQ-01"].items()}
         == {pid: [(r, w, fmt(e)) for r, w, e in b["excess_list"]] for pid, b in PB.items()})
    mark("CQ-02", X["CQ-02"] == {p: [{"sequence_no": o["sequence_no"], "operation_id": o["operation_id"], "resource_id": o["resource_id"]}
                                     for o in k.ops(p)] for p in sorted({o["product_id"] for o in k.E["Operation"]})})
    rid = P["CQ-03"]["resource_id"]
    direct = sorted({o["product_id"] for o in k.E["Operation"] if o["resource_id"] == rid})
    aff = set(direct)
    for d in direct:
        aff |= set(finished_users(d))
    mark("CQ-03", X["CQ-03"]["products_affected"] == sorted(aff) and
         X["CQ-03"]["open_orders"] == sorted(o["order_id"] for o in k.open_orders() if o["product_id"] in aff), X["CQ-03"])
    ok4 = True
    for case in X["CQ-04"]:
        c = k.pick_bom(case["product_id"], day(case["date"]), dec(case["qty"]))
        st = "SELECTED" if len(c) == 1 else ("NONE" if not c else "AMBIGUOUS")
        ok4 &= (c == case["candidates"] and st == case["status"])
    mark("CQ-04", ok4, X["CQ-04"])
    mark("CQ-05", X["CQ-05"] == {m: finished_users(m) for m in P["CQ-05"]["material_ids"]}, X["CQ-05"])
    mark("CQ-06", {m: fmt(dec(v)) for m, v in X["CQ-06"].items()} == {m: fmt(v) for m, v in sorted(R["gross"].items())})
    mark("CQ-07", {p: fmt(dec(v)) for p, v in X["CQ-07"].items()} == {p: fmt(v) for p, v in sorted(R["net"].items())})
    mark("CQ-08", {m: fmt(dec(v)) for m, v in X["CQ-08"].items()} == {m: fmt(v) for m, v in R["short"].items() if v > 0})
    mark("CQ-09", {m: [(s["supplier_id"], s["lead_time_days"], s["is_preferred"]) for s in v] for m, v in X["CQ-09"].items()}
         == {m: sorted((s["supplier_id"], s["lead_time_days"], s["is_preferred"]) for s in k.E["SupplierMaterial"] if s["material_id"] == m)
             for m in P["CQ-09"]["material_ids"]})
    n = P["CQ-10"]["min_lead_time_days"]
    longm = [s["material_id"] for s in k.E["SupplierMaterial"] if s["is_preferred"] and s["lead_time_days"] >= n]
    prod10 = defaultdict(list)
    for m in sorted(longm):
        for p in finished_users(m):
            prod10[p].append(m)
    mark("CQ-10", X["CQ-10"]["products"] == dict(prod10), X["CQ-10"])
    c11 = P["CQ-11"]
    base_lt = next(s["lead_time_days"] for s in k.E["SupplierMaterial"] if s["is_preferred"] and s["material_id"] == c11["material_id"])
    F2 = k.feasibility(R, {c11["material_id"]: base_lt + c11["delay_days"]})
    newly = sorted(o for o in F2 if not F2[o]["feasible"] and F[o]["feasible"])
    mark("CQ-11", sorted(x["order_id"] for x in X["CQ-11"]["newly_infeasible_orders"]) == newly, newly)
    sums = defaultdict(lambda: Z)
    for o in k.open_orders():
        sums[o["product_id"]] += dec(o["order_qty"])
    mark("CQ-12", {p: fmt(dec(v)) for p, v in X["CQ-12"].items()} == {p: fmt(v) for p, v in sorted(sums.items())})
    mark("CQ-13", {r["order_id"]: (r["feasible"], r["reasons"]) for r in X["CQ-13"]} == {o: (v["feasible"], v["reasons"]) for o, v in F.items()})
    mark("CQ-14", all(eq_plan(X["CQ-14"][pid], PB[pid]) for pid in PB))
    # 정상 데이터에서는 오류 관련 CQ 결과가 비어 있어야 한다 (canonical 검사와 같은 결론)
    mark("CQ-15", X["CQ-15"]["duplicate_business_ids"] == [] and _no_dup(k.E))
    mark("CQ-16", X["CQ-16"]["dangling_references"] == [] and _no_dangling(k.E))
    ok17 = all(len(v["candidates"]) == 1 and v["candidates"] == k.pick_bom(p, day(v["date"]), dec(v["qty"]))
               for p, v in X["CQ-17"]["horizon_selection"].items())
    mark("CQ-17", ok17 and X["CQ-17"]["overlapping_active_pairs"] == [])
    mark("CQ-18", X["CQ-18"]["cycles"] == [] and not find_cycles(component_graph(k.E)))
    # CQ-19: 원본 위치의 실제 셀 값을 읽어 canonical 값과 같은지 확인
    cmap = read_json(paths.column_map)
    ok19 = len(X["CQ-19"]["locations"]) == 3
    for loc in X["CQ-19"]["locations"]:
        if loc["sheet"]:
            rows = read_sheet(paths.root / "data" / loc["file"], loc["sheet"])
            hdr = rows[1] if loc["vendor"] == "B" else rows[0]
        else:
            rows = read_csv_rows(paths.root / "data" / loc["file"], next(f["encoding"] for f in cmap["files"] if f["file"] == loc["file"]))
            hdr = rows[0]
        cell = rows[loc["row_number"] - 1][hdr.index(loc["column"])]
        unit_cell = rows[loc["row_number"] - 1][hdr.index("uom" if loc["vendor"] == "A" else "단위")]
        grams = bnum(cell) * (1000 if unit_cell == "kg" else 1)
        ok19 &= str(cell) == loc["displayed_value"] and unit_cell == loc["displayed_unit"] and grams == dec(X["CQ-19"]["canonical"]["qty_per"])
    mark("CQ-19", ok19, X["CQ-19"]["locations"])
    # CQ-20: 두 업체 코드가 가리키는 제품의 속성(이름·안전재고)이 같은지 실제 파일에서 확인
    a_rows = read_csv_rows(paths.vendor_a / "products.csv", "utf-8")
    a_by = {r[0]: (r[1], dec(r[4])) for r in a_rows[1:]}
    b_rows = read_sheet(paths.vendor_b / "vendor_b_master.xlsx", "제품")
    b_by = {r[0]: (r[1], bnum(r[4])) for r in b_rows[2:]}
    mark("CQ-20", all(a_by[v["vendor_a_code"]] == b_by[v["vendor_b_code"]] for v in X["CQ-20"].values()) and len(X["CQ-20"]) == 5)
    unverified = {c: "미검증" for c in [f"CQ-{i:02d}" for i in range(1, 21)] if c not in status}
    return {"status": {**status, **unverified}}


def eq_plan(a: dict, b: dict) -> bool:
    c = a["costs_krw"]
    f = lambda x: fmt(dec(x))
    return (f(c["material"]) == fmt(b["material"]) and f(c["regular_production"]) == fmt(b["regular_production"])
            and f(c["overtime"]) == fmt(b["overtime"]) and f(c["outsourcing"]) == fmt(b["outsourcing"])
            and f(a["worked_overtime_h"]) == fmt(b["worked_overtime_h"]) and f(a["outsourced_qty_total"]) == fmt(b["outsourced_qty"])
            and f(a["capacity_excess_total_h"]) == fmt(b["excess_total"]) and f(a["unproduced_qty_total"]) == fmt(b["unproduced"])
            and sorted(a["late_orders"]) == b["late"])


def _no_dup(E):
    return all(len({tuple(r[k] for k in keys) for r in E[e]}) == len(E[e]) for e, keys in PK.items())


def _no_dangling(E):
    prods = {p["product_id"] for p in E["Product"]}
    mats = {m["material_id"] for m in E["Material"]}
    res = {r["resource_id"] for r in E["Resource"]}
    return (all((l["component_material_id"] or l["component_product_id"]) in prods | mats for l in E["BOMLine"])
            and all(o["product_id"] in prods for o in E["Order"]) and all(o["resource_id"] in res for o in E["Operation"]))


def run(paths: Paths) -> dict:
    gt = read_json(paths.ground_truth)
    C = Checks()
    check_canonical(gt, paths, C)
    check_vendors(gt, paths, C)
    inj = check_injection(gt, paths, C)
    cq = check_expected(gt, paths, C)
    report = {"implementation": "scripts/e01_verify.py (구현 B, 생성·정답 모듈 비의존)", "summary": C.summary(),
              "injection": inj, "cq_status": cq["status"], "checks": C.items}
    write_json(paths.results / "validation_report.json", json.loads(json.dumps(report, default=str, ensure_ascii=False)))
    return report


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out-root", type=Path, default=None)
    args = ap.parse_args(argv)
    rep = run(Paths(args.out_root))
    print(json.dumps(rep["summary"], ensure_ascii=False))
    for c in rep["checks"]:
        if not c["passed"]:
            print("FAIL", c["section"], c["check"], str(c["detail"])[:400])
    return 0 if rep["summary"]["failed"] == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
