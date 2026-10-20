"""E01 오류 주입: 정상 Vendor A/B 파일을 data/invalid/에 복사한 뒤 오류를 넣는다.

    .\\.venv\\Scripts\\python.exe scripts\\e01_inject_errors.py [--out-root DIR]

- 정상 데이터셋(data/vendor_a, data/vendor_b)은 바꾸지 않는다.
- 각 오류의 고유 ID·파일·행·필드·유형·바꾼 값·기대 탐지 결과를 data/canonical/injected_errors_v1.json에 기록한다.
- 주입이 실제로 일어났는지는 e01_verify.py가 정상본과 주입본을 독립적으로 비교해 확인한다(이 모듈을 import하지 않음).
"""

from __future__ import annotations

import os

os.environ.setdefault("OPENPYXL_LXML", "False")  # E01 xlsx 바이트 재현 (e01_generate.py 주석 참고)

import argparse
import csv
import io
import shutil
import sys
from pathlib import Path

import openpyxl

from e01_common import Paths, normalize_xlsx, write_json


class CsvFile:
    def __init__(self, path: Path, encoding: str, newline: str) -> None:
        self.path, self.encoding, self.newline = path, encoding, newline
        self.rows = list(csv.reader(io.StringIO(path.read_bytes().decode(encoding), newline="")))
        self.header = self.rows[0]

    def find(self, **match) -> int:
        idx = [n for n, r in enumerate(self.rows[1:], start=2)
               if all(r[self.header.index(k)] == v for k, v in match.items())]
        if len(idx) != 1:
            raise ValueError(f"{self.path.name}: {match} → {len(idx)}행")
        return idx[0]

    def get(self, row: int, col: str) -> str:
        return self.rows[row - 1][self.header.index(col)]

    def set(self, row: int, col: str, value: str) -> str:
        old = self.get(row, col)
        self.rows[row - 1][self.header.index(col)] = value
        return old

    def append(self, values: list[str]) -> int:
        self.rows.append(values)
        return len(self.rows)

    def save(self) -> None:
        buf = io.StringIO()
        csv.writer(buf, lineterminator=self.newline).writerows(self.rows)
        self.path.write_bytes(buf.getvalue().encode(self.encoding))


def _xlsx_find(ws, header_row: int, **match) -> int:
    hdr = [c.value for c in ws[header_row]]
    hits = [r for r in range(header_row + 1, ws.max_row + 1)
            if all(str(ws.cell(r, hdr.index(k) + 1).value) == v for k, v in match.items())]
    if len(hits) != 1:
        raise ValueError(f"{ws.title}: {match} → {len(hits)}행")
    return hits[0]


def inject(paths: Paths) -> list[dict]:
    inv_a, inv_b = paths.invalid / "vendor_a", paths.invalid / "vendor_b"
    for src, dst in ((paths.vendor_a, inv_a), (paths.vendor_b, inv_b)):
        if dst.exists():
            shutil.rmtree(dst)
        shutil.copytree(src, dst)
    errors: list[dict] = []

    def add(eid, code, vendor, file, sheet, row, field, canonical_field, entity, key, desc, changes, rule, expected, derived=None):
        errors.append({"error_id": eid, "error_code": code, "vendor": vendor, "file": f"invalid/{file}", "clean_file": file,
                       "sheet": sheet, "row_number": row, "field": field, "canonical_field": canonical_field,
                       "canonical_entity": entity, "business_key": key, "description": desc, "changes": changes,
                       "expected_detection": {"rule": rule, "result": expected}, "derived_effects": derived or []})

    # ---- Vendor A CSV ----
    orders = CsvFile(inv_a / "orders.csv", "utf-8", "\n")
    r = orders.find(order_no="SO260004")
    old = orders.set(r, "qty", "")
    add("E01-ERR-001", "ERR-MISSING", "A", "vendor_a/orders.csv", "", r, "qty", "order_qty", "Order", "O004",
        "주문 수량 필수값 누락", [{"row": r, "column": "qty", "old": old, "new": ""}], "ERR-MISSING (필수 필드)",
        "Order O004 레코드의 order_qty 필드 누락으로 보고")
    r = orders.find(order_no="SO260007")
    old = orders.set(r, "due_date", "2026-02-30")
    add("E01-ERR-008", "ERR-DATE", "A", "vendor_a/orders.csv", "", r, "due_date", "due_date", "Order", "O007",
        "존재하지 않는 날짜", [{"row": r, "column": "due_date", "old": old, "new": "2026-02-30"}], "QR-04 ①",
        "Order O007 due_date 형식·날짜 오류로 보고", ["O007은 납기 판정 불가(unknown)"])
    orders.save()

    mats = CsvFile(inv_a / "materials.csv", "utf-8-sig", "\n")
    orig = mats.find(material_code="000105")
    r = mats.append(["000105", "PCB 기판(중복 등록)", "EA", "4700", "50"])
    add("E01-ERR-002", "ERR-DUPID", "A", "vendor_a/materials.csv", "", r, "material_code", "material_id", "Material", "M005",
        "같은 자재 코드를 다른 이름·단가로 중복 등록", [{"row": r, "append": True}], "QR-01",
        f"material_code 000105 레코드 2건(행 {orig}, {r}) 모두 중복 ID로 보고 (필드 *)",
        ["Vendor A BOM의 000105 참조(BOM-SF001-V1 line 2)가 어느 자재인지 모호해짐 → 파생 보고"])
    mats.save()

    bom = CsvFile(inv_a / "bom.csv", "utf-8", "\n")
    r = bom.find(bom_code="BM-0002001-1", line_no="3")
    old = bom.set(r, "uom", "EA")
    add("E01-ERR-004", "ERR-UNIT", "A", "vendor_a/bom.csv", "", r, "uom", "unit", "BOMLine", "BOM-SF001-V1|3",
        "kg 자재(ABS 수지)에 EA 단위", [{"row": r, "column": "uom", "old": old, "new": "EA"}], "QR-03 ②",
        "BOMLine BOM-SF001-V1 line 3의 unit 차원 불일치(MASS 자재에 COUNT)로 보고")
    r = bom.find(bom_code="BM-0001002-1", line_no="3")
    old = bom.set(r, "bom_version", "V2")
    add("E01-ERR-007", "ERR-IDVER", "A", "vendor_a/bom.csv", "", r, "bom_version", "bom_version", "BOMLine", "BOM-FG002-V1|3",
        "같은 bom_code의 다른 행과 버전 표기가 다름 (ID·버전 불일치)", [{"row": r, "column": "bom_version", "old": old, "new": "V2"}],
        "BOM 헤더 일관성 (ERR-IDVER)", "bom_code BM-0001002-1의 (product_code, bom_version)이 행마다 다름으로 보고")
    hdr = bom.header
    vals = {"bom_code": "BM-0001002-2", "product_code": "0001002", "bom_version": "V2", "status": "APPROVED", "active": "Y",
            "valid_from": "2026-11-01", "valid_to": "", "min_qty": "0", "max_qty": "", "site": "", "base_qty": "1",
            "line_no": "1", "component_type": "PRD", "component_code": "0002001", "qty_per": "2", "uom": "EA", "scrap_rate": "0"}
    r = bom.append([vals[h] for h in hdr])
    add("E01-ERR-011", "ERR-BOMVER", "A", "vendor_a/bom.csv", "", r, "valid_from", "valid_from", "BOM", "FG002 V2",
        "FG002에 V1과 기간·수량·사이트 조건이 모두 겹치는 활성 V2 추가", [{"row": r, "append": True}], "QR-06 ①",
        "FG002의 활성 버전 V1(BM-0001002-1)과 V2(BM-0001002-2) 유효 조건 중복(기간·수량·사이트)으로 보고",
        ["2026-11-01 이후 FG002 BOM 선택 모호 → FG002 계산 중단"])
    bom.save()

    inv = CsvFile(inv_a / "inventory.csv", "utf-8", "\n")
    r = inv.find(item_code="0001002")
    old = inv.set(r, "allocated", "80")
    add("E01-ERR-012", "ERR-INVQTY", "A", "vendor_a/inventory.csv", "", r, "allocated", "allocated_qty", "Inventory",
        "PRODUCT|FG002|2026-11-15", "할당 수량이 보유 수량(60)보다 큼", [{"row": r, "column": "allocated", "old": old, "new": "80"}],
        "QR-07 ②", "Inventory FG002의 allocated_qty > on_hand_qty로 보고")
    inv.save()

    rt = CsvFile(inv_a / "routing.csv", "utf-8", "\n")
    r = rt.find(operation_code="OP-FG003-20")
    old = rt.set(r, "resource_code", "R09")
    add("E01-ERR-013", "ERR-RESREL", "A", "vendor_a/routing.csv", "", r, "resource_code", "resource_id", "Operation", "OP-FG003-20",
        "존재하지 않는 자원 참조", [{"row": r, "column": "resource_code", "old": old, "new": "R09"}], "QR-08 ①",
        "Operation OP-FG003-20의 resource_id가 Resource에 없음으로 보고", ["FG003 공정 경로 불완전 → 능력 계산 불가"])
    rt.save()

    # ---- Vendor B CSV (CP949) ----
    bo = CsvFile(inv_b / "vendor_b_orders.csv", "cp949", "\r\n")
    r = bo.find(주문번호="주문-012")
    old = bo.set(r, "납기일자", "2026.10.01")
    add("E01-ERR-009", "ERR-DATE", "B", "vendor_b/vendor_b_orders.csv", "", r, "납기일자", "due_date", "Order", "O012",
        "납기일이 주문일(2026.11.11)보다 빠름", [{"row": r, "column": "납기일자", "old": old, "new": "2026.10.01"}], "QR-04 ②",
        "Order O012 due_date < order_date로 보고")
    bo.save()
    bi = CsvFile(inv_b / "vendor_b_inventory.csv", "cp949", "\r\n")
    r = bi.find(품목코드="MAT-01")
    old = bi.set(r, "현재고", "8OO,000")
    add("E01-ERR-005", "ERR-NUMFMT", "B", "vendor_b/vendor_b_inventory.csv", "", r, "현재고", "on_hand_qty", "Inventory",
        "MATERIAL|M001|2026-11-15", "숫자 0 대신 영문자 O", [{"row": r, "column": "현재고", "old": old, "new": "8OO,000"}],
        "ERR-NUMFMT", "Inventory M001 on_hand_qty 숫자 형식 오류로 보고")
    bi.save()

    # ---- Vendor B XLSX ----
    xp = inv_b / "vendor_b_master.xlsx"
    wb = openpyxl.load_workbook(xp)
    ws = wb["BOM"]
    r = _xlsx_find(ws, 2, BOM번호="B-BOM-06", 순번="2")
    col = [c.value for c in ws[2]].index("자품목코드") + 1
    old = ws.cell(r, col).value
    ws.cell(r, col).value = "MAT-99"
    add("E01-ERR-003", "ERR-BADREF", "B", "vendor_b/vendor_b_master.xlsx", "BOM", r, "자품목코드", "component_material_id",
        "BOMLine", "BOM-SF001-V1|2", "존재하지 않는 자재 코드 참조", [{"row": r, "column": "자품목코드", "old": old, "new": "MAT-99"}],
        "QR-02", "BOMLine BOM-SF001-V1 line 2의 component_material_id가 Material에 없음으로 보고",
        ["Vendor B 기준 SF001의 PCB 기판(M005) 소요가 빠짐"])
    hdr = [c.value for c in ws[2]]
    vals = {"BOM번호": "B-BOM-07", "모품목코드": "MD-HOUS", "버전": "V1", "상태": "승인", "활성": "예", "시작일": "2026.01.01",
            "종료일": None, "최소수량": "0", "최대수량": None, "사이트": None, "기준수량": "1", "순번": 4, "자품목구분": "제품",
            "자품목코드": "BL-C", "소요량": "1", "단위": "개", "손실률(%)": "0"}
    r = ws.max_row + 1
    for c, h in enumerate(hdr, start=1):
        ws.cell(r, c).value = vals[h]
        ws.cell(r, c).number_format = "@"
    add("E01-ERR-010", "ERR-BOMCYCLE", "B", "vendor_b/vendor_b_master.xlsx", "BOM", r, "자품목코드", "component_product_id",
        "BOMLine", "BOM-SF002-V1|4", "반제품 SF002의 구성품으로 완제품 FG003 추가 → FG003→SF002→FG003 순환",
        [{"row": r, "append": True}], "QR-05", "순환 경로 FG003 → SF002 → FG003의 BOMLine 레코드로 보고",
        ["Vendor B 기준 FG003·SF002 전개 불가"])
    ws = wb["자재"]
    changes = []
    shifted = {3: "단가(원/단위)", 4: "안전재고", 5: None}
    for c, v in shifted.items():
        changes.append({"row": 2, "column_index": c, "old": ws.cell(2, c).value, "new": v})
        ws.cell(2, c).value = v
    add("E01-ERR-006", "ERR-TABLE", "B", "vendor_b/vendor_b_master.xlsx", "자재", 2, "*", "*", "Material", "*",
        "2행 헤더에서 '단위' 칸이 빠지고 뒤 열 이름이 한 칸씩 당겨짐 (헤더-데이터 열 어긋남)", changes,
        "표 구조 보존 (ERR-TABLE)", "시트 '자재'의 헤더 열 구성 불일치로 보고 (레코드 *, 필드 *)",
        ["자재 시트 10행 전체가 잘못된 열로 해석될 위험"])
    wb.save(xp)
    normalize_xlsx(xp)
    errors.sort(key=lambda e: e["error_id"])
    write_json(paths.injected, {
        "dataset": "e01_manufacturing_v1 invalid copy",
        "clean_dirs": ["data/vendor_a", "data/vendor_b"], "invalid_dirs": ["data/invalid/vendor_a", "data/invalid/vendor_b"],
        "row_number_rule": "파일·시트의 물리적 행 번호(1부터, 헤더 포함)",
        "error_count": len(errors), "error_codes": sorted({e["error_code"] for e in errors}),
        "errors": errors,
    })
    return errors


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out-root", type=Path, default=None)
    args = ap.parse_args(argv)
    errs = inject(Paths(args.out_root))
    print(len(errs), "errors:", sorted({e["error_code"] for e in errs}))
    return 0


if __name__ == "__main__":
    sys.exit(main())
