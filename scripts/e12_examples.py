"""E12 형식별 실제 사례 — ① 원본 ② 기존 파서 출력 ③ OfficeCLI 출력 ④ 비교 ⑤ 결론.

    .\\.venv\\Scripts\\python.exe scripts\\e12_examples.py

- 모든 값은 이 스크립트 실행 중 실제로 읽은 것이다: 원본 XML 발췌(zip 안 XML), openpyxl 실행 결과, E03 v2 IR(e12_run.py 출력 파일),
  OfficeCLI `get <경로> --json` 원시 노드(임시 복사본에 실행)와 변환된 IR. 결론 문장만 사람이 쓴 해설이다.
- 공개 사례(합성 문서): results/E12_officecli/examples.json
- 외부 문서 사례(공개 문서 본문 포함, D-025): data/raw/e12_officecli/examples_external.json (Git 제외)
- e12_run.py 실행 후에 돌린다 (outputs/synthetic/*.json 필요).
"""

from __future__ import annotations

import json
import re
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import e12_officecli as O  # noqa: E402

OUT = ROOT / "results" / "E12_officecli"
SYN = OUT / "outputs" / "synthetic"
RAW = ROOT / "data" / "raw" / "e12_officecli"
EXT_TRUTH = ROOT / "data" / "canonical" / "external" / "e03_external_truth_v1.json"
X12 = ROOT / "data" / "e12" / "X12-01_orders.xlsx"
X12NC = ROOT / "data" / "e12" / "X12-02_orders_no_formula_cache.xlsx"
XTRUTH = json.loads((ROOT / "data" / "canonical" / "e12" / "xlsx_truth_v1.json").read_text(encoding="utf-8"))


def xml_part(path: Path, name: str) -> str:
    with zipfile.ZipFile(path) as z:
        return z.read(name).decode("utf-8")


def excerpt(xml: str, pattern: str, limit=600) -> str | None:
    m = re.search(pattern, xml, re.S)
    return m.group(0)[:limit] if m else None


def oc_get(src: Path, path: str, depth: int = 1) -> dict:
    r = O.run(["get", "{file}", path, "--depth", str(depth), "--json"], src)
    res = O.results(r["data"].get("data"))
    node = res[0] if res else {"error": r["data"].get("error")}
    return {"command": f"officecli get <file> {path} --depth {depth} --json", "returncode": r["returncode"], "source_unchanged": r["source_unchanged"], "node": prune(node)}


def prune(n, keep_children=True):
    """원시 노드에서 서식 출처(effective.*) 같은 잡음만 뺀다 — 값·병합·경로 필드는 그대로."""
    if isinstance(n, dict):
        out = {}
        for k, v in n.items():
            if k == "format" and isinstance(v, dict):
                out[k] = {a: b for a, b in v.items() if not a.startswith("effective.")}
            elif k == "children":
                out[k] = [prune(c) for c in v] if keep_children else f"{len(v)}개 생략"
            else:
                out[k] = v
        return out
    return n


def sheet_xml(path: Path, sheet_index: int) -> str:
    return xml_part(path, f"xl/worksheets/sheet{sheet_index}.xml")


def openpyxl_cell(path: Path, sheet: str, ref: str) -> dict:
    import openpyxl
    wf, wv = openpyxl.load_workbook(path, data_only=False), openpyxl.load_workbook(path, data_only=True)
    c, v = wf[sheet][ref], wv[sheet][ref]
    return {"code": f"load_workbook(data_only=False/True)['{sheet}']['{ref}']", "value(data_only=False)": repr(c.value), "value(data_only=True)": repr(v.value),
            "data_type": c.data_type, "number_format": c.number_format, "is_date": c.is_date}


def truth_cell(sheet, ref):
    return next(t for t in XTRUTH["cells"] if t["sheet"] == sheet and t["cell"] == ref)


# ------------------------------------------------------------------ XLSX
def xlsx_examples() -> list[dict]:
    import openpyxl
    ex = []
    # 1. 문자열로 저장된 숫자
    ex.append({"id": "XLSX-1", "title": "문자열로 저장된 숫자 (Orders!C3 = \"30\")", "file": "data/e12/X12-01_orders.xlsx",
               "original": {"xml": excerpt(sheet_xml(X12, 1), r'<c r="C3"[^>]*>.*?</c>'), "truth": truth_cell("Orders", "C3")},
               "baseline": openpyxl_cell(X12, "Orders", "C3"), "officecli": oc_get(X12, "/Orders/C3")})
    # 2. 수식 캐시가 없는 파일의 수식 셀
    ex.append({"id": "XLSX-2", "title": "수식 캐시값이 0으로 남은 파일의 수식 셀 (X12-02 Orders!F2 = C2*E2)", "file": "data/e12/X12-02_orders_no_formula_cache.xlsx",
               "original": {"xml": excerpt(sheet_xml(X12NC, 1), r'<c r="F2"[^>]*>.*?</c>'), "truth": truth_cell("Orders", "F2")},
               "baseline": openpyxl_cell(X12NC, "Orders", "F2"), "officecli": oc_get(X12NC, "/Orders/F2")})
    # 3. 병합 셀
    wb = openpyxl.load_workbook(X12)
    ex.append({"id": "XLSX-3", "title": "병합 셀 (Orders!A8:E8, BOM!E3:E4 세로 병합)", "file": "data/e12/X12-01_orders.xlsx",
               "original": {"xml": excerpt(sheet_xml(X12, 1), r"<mergeCells.*?</mergeCells>"), "xml_bom": excerpt(sheet_xml(X12, 2), r"<mergeCells.*?</mergeCells>"),
                            "truth": [m for m in XTRUTH["merges"]]},
               "baseline": {"code": "ws.merged_cells.ranges", "Orders": [str(r) for r in wb["Orders"].merged_cells.ranges], "BOM": [str(r) for r in wb["BOM"].merged_cells.ranges],
                            "Orders!B8 (덮인 셀) value": repr(wb["Orders"]["B8"].value)},
               "officecli": {"A8": oc_get(X12, "/Orders/A8"), "B8": oc_get(X12, "/Orders/B8"), "E4": oc_get(X12, "/BOM/E4")}})
    # 4. 숨긴 행·열
    ws = wb["Orders"]
    cols = O.run(["query", "{file}", "column", "--json"], X12)
    ex.append({"id": "XLSX-4", "title": "숨긴 행·열 (Orders 5행, I열)", "file": "data/e12/X12-01_orders.xlsx",
               "original": {"xml_row": excerpt(sheet_xml(X12, 1), r'<row r="5"[^>]*>'), "xml_cols": excerpt(sheet_xml(X12, 1), r"<cols>.*?</cols>"),
                            "truth": {"hidden_rows": XTRUTH["hidden_rows"], "hidden_cols": XTRUTH["hidden_cols"]}},
               "baseline": {"code": "ws.row_dimensions[5].hidden / ws.column_dimensions['I'].hidden", "row5_hidden": ws.row_dimensions[5].hidden,
                            "colI_hidden": ws.column_dimensions["I"].hidden},
               "officecli": {"row5": oc_get(X12, "/Orders/row[5]", 0),
                             "query column": {"command": "officecli query <file> column --json", "returncode": cols["returncode"],
                                              "hidden": [prune(n, False) for n in O.results(cols["data"].get("data")) if n.get("format", {}).get("hidden")]}}})
    # 5. 수식 오류값
    ex.append({"id": "XLSX-5", "title": "잘못된 입력값이 만든 수식 오류 (Orders!C4 = \"1O50\", F4 = C4*E4)", "file": "data/e12/X12-01_orders.xlsx",
               "original": {"xml": excerpt(sheet_xml(X12, 1), r'<c r="F4"[^>]*>.*?</c>'), "truth": [truth_cell("Orders", "C4"), truth_cell("Orders", "F4")]},
               "baseline": {"C4": openpyxl_cell(X12, "Orders", "C4"), "F4": openpyxl_cell(X12, "Orders", "F4")},
               "officecli": {"C4": oc_get(X12, "/Orders/C4"), "F4": oc_get(X12, "/Orders/F4")}})
    r = json.loads((OUT / "results.json").read_text(encoding="utf-8"))["xlsx"]
    concl = {
        "XLSX-1": "두 도구 모두 문자열로 보존한다 (openpyxl data_type 's', OfficeCLI type SharedString). 차이 없음.",
        "XLSX-2": f"openpyxl은 파일의 캐시값 0을 그대로 돌려준다 (X12-02 값 {r['X12-02_no_formula_cache']['openpyxl']['metrics']['value']}). "
                  f"OfficeCLI는 cachedValue 0과 자체 계산한 computedValue를 함께 준다 (computedValue 사용 시 {r['X12-02_no_formula_cache']['officecli']['metrics']['value']}). "
                  "다른 도구로 저장되어 캐시가 없는 파일에서만 의미 있는 차이다. 계산 엔진 결과는 Excel 계산과 다를 수 있어 검증 대상이다.",
        "XLSX-3": "병합 범위는 같다. OfficeCLI는 덮인 셀마다 merge 범위를 붙여 주고 앵커 표시(mergeAnchor)를 준다. openpyxl은 범위 목록으로 준다. 정확도 차이 없음.",
        "XLSX-4": "숨긴 행·열 모두 같게 찾는다. OfficeCLI는 열 정보를 따로 query column으로 받아야 한다 (호출 1회 추가, batch에 포함).",
        "XLSX-5": "두 도구 모두 C4를 문자열 \"1O50\"으로, F4 캐시값을 #VALUE!로 준다. 잘못된 입력 판정은 어느 도구도 하지 않으며 E05 검증 단계의 몫이다.",
    }
    for e in ex:
        e["comparison"] = "아래 baseline/officecli 필드 대조"
        e["conclusion"] = concl[e["id"]]
    return ex


# ------------------------------------------------------------------ DOCX·PPTX (합성)
def ir(doc, kind):
    return json.loads((SYN / f"{doc}__{kind}.json").read_text(encoding="utf-8"))


def cell(irx, tid, row, col):
    t = next(t for t in irx["tables"] if t["table_id"] == tid)
    c = next((c for c in t["cells"] if c["row"] == row and c["col"] == col), None)
    return {"table": tid, "n_rows": t["n_rows"], "n_cols": t["n_cols"], "cell": {k: c[k] for k in ("row", "col", "rowspan", "colspan", "text")} if c else None,
            "location": c["location"] if c else None}


def blocks(irx, typ):
    return [{"type": b["type"], "text": b["text"], "location": b["location"]} for b in irx["blocks"] if b["type"] == typ]


def docx_tc_xml(path: Path, t_index: int, r_index: int, c_index: int) -> str | None:
    xml = xml_part(path, "word/document.xml")
    tbls = re.findall(r"<w:tbl>.*?</w:tbl>", xml, re.S)
    if t_index >= len(tbls):
        return None
    rows = re.findall(r"<w:tr[ >].*?</w:tr>", tbls[t_index], re.S)
    tcs = re.findall(r"<w:tc>.*?</w:tc>", rows[r_index], re.S)
    return re.sub(r"<w:rPr>.*?</w:rPr>", "", tcs[c_index], flags=re.S)[:700]


def docx_examples() -> list[dict]:
    d1, d2 = ROOT / "data/documents/D-01_bom_spec.docx", ROOT / "data/documents/D-02_order_inventory.docx"
    b1, o1, b2, o2 = ir("D-01", "v2"), ir("D-01", "officecli"), ir("D-02", "v2"), ir("D-02", "officecli")
    ex = [
        {"id": "DOCX-1", "title": "병합 머리글 셀 (D-01 첫 표 1행: 세로 병합 '순번'·가로 병합 '구성품')", "file": "data/documents/D-01_bom_spec.docx",
         "original": {"xml_tc_1_1": docx_tc_xml(d1, 0, 0, 0), "xml_tc_1_2": docx_tc_xml(d1, 0, 0, 1), "xml_tc_2_1": docx_tc_xml(d1, 0, 1, 0)},
         "baseline": [cell(b1, "t001", 1, 1), cell(b1, "t001", 1, 2)], "officecli_ir": [cell(o1, "t001", 1, 1), cell(o1, "t001", 1, 2)],
         "officecli": {"tc_1_1": oc_get(d1, "/body/tbl[1]/tr[1]/tc[1]", 0), "tc_1_2": oc_get(d1, "/body/tbl[1]/tr[1]/tc[2]", 0), "tc_2_1": oc_get(d1, "/body/tbl[1]/tr[2]/tc[1]", 0)},
         "conclusion": "OfficeCLI는 colspan(gridSpan)과 vmerge=restart/continue를 원시 값으로 준다. rowspan은 주지 않으므로 아래 행의 continue를 세는 변환이 필요하다 "
                       "(e12_officecli._grid). 변환 후 v2와 같은 병합 결과."},
        {"id": "DOCX-2", "title": "바닥글 쪽 번호 필드 (D-01 footer)", "file": "data/documents/D-01_bom_spec.docx",
         "original": {"xml_footer": excerpt(xml_part(d1, "word/footer1.xml") if "word/footer1.xml" in zipfile.ZipFile(d1).namelist() else "", r"<w:p[ >].*?</w:p>", 900)},
         "baseline": blocks(b1, "footer"), "officecli_ir": blocks(o1, "footer"), "officecli": oc_get(d1, "/footer[1]", 3),
         "conclusion": "OfficeCLI는 fldSimple 필드 결과(쪽 번호 캐시값 '1')를 문단 text에서는 빼고 footer 부분(part) text에만 넣는다. "
                       "첫 변환 코드는 문단 text만 써서 빈 값이 되었고(변환기 문제), 문단이 하나인 머리글·바닥글은 part text를 쓰도록 고친 뒤 v2와 같다. "
                       "본문 문단 안의 필드(외부 X-D04 '1 of 2'의 NUMPAGES)는 대신 쓸 상위 text가 없어 여전히 빠진다."},
        {"id": "DOCX-3", "title": "주문 표의 수량·단위 셀 (D-02 첫 표)", "file": "data/documents/D-02_order_inventory.docx",
         "original": {"xml_tc": docx_tc_xml(d2, 0, 2, 3)}, "baseline": [cell(b2, "t001", 3, 4)], "officecli_ir": [cell(o2, "t001", 3, 4)],
         "officecli": oc_get(d2, "/body/tbl[1]/tr[3]/tc[4]", 0),
         "conclusion": "셀 텍스트·위치 같음. OfficeCLI 위치는 /body/tbl[1]/tr[3]/tc[4] 경로로, v2는 body_index·row·col로 준다. 데이터 형식(숫자·단위) 해석은 두 도구 모두 하지 않는다."},
    ]
    for e in ex:
        e["comparison"] = "baseline(E03 v2 IR) / officecli_ir(OfficeCLI → 같은 IR 변환) 대조"
    return ex


def pptx_xml_tc(path: Path, slide: int, t_row: int, n: int) -> list[str]:
    xml = xml_part(path, f"ppt/slides/slide{slide}.xml")
    rows = re.findall(r"<a:tr[ >].*?</a:tr>", xml, re.S)
    if t_row >= len(rows):
        return []
    tcs = re.findall(r"<a:tc[ >/].*?</a:tc>|<a:tc [^>]*/>", rows[t_row], re.S)
    return [re.sub(r"<a:rPr.*?(/>|</a:rPr>)", "", x, flags=re.S)[:400] for x in tcs[:n]]


def pptx_examples() -> list[dict]:
    t1, t2 = ROOT / "data/documents/T-01_plan_comparison.pptx", ROOT / "data/documents/T-02_bom_spec.pptx"
    b1, o1, b2, o2 = ir("T-01", "v2"), ir("T-01", "officecli"), ir("T-02", "v2"), ir("T-02", "officecli")
    tl = lambda x: [{k: t.get(k) for k in ("table_id", "n_rows", "n_cols", "continuation", "repeated_header_rows")} | {"slide_index": t["location"].get("slide_index")} for t in x["tables"][:2]]
    ex = [
        {"id": "PPTX-1", "title": "병합 머리글 셀 (T-02 슬라이드 2 표 1행)", "file": "data/documents/T-02_bom_spec.pptx",
         "original": {"xml_row1_tc1_3": pptx_xml_tc(t2, 2, 0, 3)},
         "baseline": [cell(b2, "t001", 1, c) for c in (1, 2, 4, 6)], "officecli_ir": [cell(o2, "t001", 1, c) for c in (1, 2, 4, 6)],
         "officecli": {"tc_1_2": oc_get(t2, "/slide[2]/table[1]/tr[1]/tc[2]", 0), "tc_1_3": oc_get(t2, "/slide[2]/table[1]/tr[1]/tc[3]", 0)},
         "conclusion": "OfficeCLI는 앵커에 colspan/rowspan, 덮인 셀에 hmerge/vmerge를 준다. 덮인 셀도 tc로 남으므로 tc 하나가 격자 1열이다. "
                       "첫 변환 코드는 앵커 뒤에서 colspan만큼 건너뛰어 열 수를 8로 잘못 만들었고(변환기 버그, 수정 후 재실행), 수정 후 v2와 같다."},
        {"id": "PPTX-2", "title": "발표자 노트 (T-01)", "file": "data/documents/T-01_plan_comparison.pptx",
         "original": {"xml_notes": excerpt(xml_part(t1, next(n for n in zipfile.ZipFile(t1).namelist() if n.startswith("ppt/notesSlides/") and n.endswith(".xml"))),
                                           r"<p:txBody>.*?</p:txBody>", 900)},
         "baseline": blocks(b1, "speaker_notes"), "officecli_ir": blocks(o1, "speaker_notes"),
         "officecli": {"command": "officecli query <file> notes --json", **{k: v for k, v in O.run(["query", "{file}", "notes", "--json"], t1).items() if k in ("returncode", "source_unchanged")},
                       "results": [prune(n) for n in O.results(O.run(["query", "{file}", "notes", "--json"], t1)["data"].get("data"))]},
         "conclusion": "노트 텍스트·슬라이드 번호 같음. OfficeCLI는 query notes 한 번으로 모든 슬라이드 노트를 준다."},
        {"id": "PPTX-3", "title": "슬라이드에 걸쳐 이어지는 표 (T-01 슬라이드 2→3)", "file": "data/documents/T-01_plan_comparison.pptx",
         "original": {"xml_slide2_row1": pptx_xml_tc(t1, 2, 0, 5), "xml_slide3_row1": pptx_xml_tc(t1, 3, 0, 5)},
         "baseline": tl(b1), "officecli_ir": tl(o1),
         "conclusion": "두 도구 모두 표 조각 연결 기능은 없다. 연결은 E03 공통 후처리(header_repeat_rule)를 같은 방식으로 적용한 결과다. "
                       "첫 실행에는 OfficeCLI 쪽에 이 규칙을 빠뜨려 0/1이 나왔고, 규칙을 같게 적용한 뒤 1/1이 되었다."},
    ]
    for e in ex:
        e["comparison"] = "baseline(E03 v2 IR) / officecli_ir(OfficeCLI → 같은 IR 변환) 대조"
    return ex


# ------------------------------------------------------------------ 외부 문서 (로컬 전용)
def external_examples() -> list[dict]:
    xt = {x["doc_id"]: x for x in json.loads(EXT_TRUTH.read_text(encoding="utf-8"))["documents"]}
    out = []

    def src(did):
        d = xt[did]
        return ROOT / f"data/raw/external_documents/{d['format']}/{d['file']}"

    def raw_ir(did, kind):
        p = ROOT / "data/raw/e12_officecli/outputs" / f"{did}__{kind}.json"
        return json.loads(p.read_text(encoding="utf-8")) if p.exists() else None

    cases = [("X-HO-T03", "pptx", "SmartArt 노드 텍스트", "/slide[1]", 2), ("X-HO-D03", "docx", "문단 안 콘텐츠 컨트롤(인라인 sdt)", "/body/p[5]", 1),
             ("X-HO-D01", "docx", "바닥글 PAGE 필드", "/footer[1]", 3), ("X-D03", "docx", "콘텐츠 컨트롤 안의 표", "/body/sdt[4]", 2)]
    for did, fmt, title, path, depth in cases:
        b, o = raw_ir(did, "v2"), raw_ir(did, "officecli")
        out.append({"id": did, "split": xt[did]["split"], "title": title, "truth_status": "PROVISIONAL (잠정 정답)",
                    "original_truth": xt[did]["truth"], "baseline_blocks": (b or {}).get("blocks", [])[:60], "officecli_blocks": (o or {}).get("blocks", [])[:60],
                    "officecli_diagnostics": (o or {}).get("diagnostics"), "officecli": oc_get(src(did), path, depth)})
    return out


def main() -> int:
    pub = {"note": "모든 출력은 실제 실행 결과. 'conclusion'만 해설. 외부 문서 사례는 data/raw/e12_officecli/examples_external.json (로컬 전용)",
           "xlsx": xlsx_examples(), "docx": docx_examples(), "pptx": pptx_examples()}
    (OUT / "examples.json").write_text(json.dumps(pub, ensure_ascii=False, indent=1, default=str) + "\n", encoding="utf-8", newline="\n")
    RAW.mkdir(parents=True, exist_ok=True)
    (RAW / "examples_external.json").write_text(json.dumps(external_examples(), ensure_ascii=False, indent=1, default=str) + "\n", encoding="utf-8", newline="\n")
    print({k: len(v) for k, v in pub.items() if isinstance(v, list)})
    return 0


if __name__ == "__main__":
    sys.exit(main())
